"""Tests for tools/ingest_expression.py, run against small synthetic inputs."""

import copy
import csv
import gzip
import hashlib
import io
import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "tools"))

import ingest_expression as ingest  # noqa: E402

CROSSWALK_HEADER = [
    "subject_locus_tag", "relationship", "object_namespace", "object_id", "seqid",
    "start", "end", "strand", "mapping_ambiguity", "source", "evidence", "mapping_method",
]
CONDITIONS = (
    "CONDITION SET: control | SAMPLES: GSM1 ;; "
    'temperature = 37°C [GEO growth protocol; quote: "Grown at 37℃"] ;; '
    "light_intensity = not reported ;; "
    'co2 = ambient [paper Methods; quote: "ambient air"] ;; '
    'medium = BG-11 [GEO growth protocol; quote: "in BG11 liquid medium"] ;; '
    'culture_format = liquid [GEO growth protocol; quote: "liquid medium"] ;; '
    'growth_phase = day 4 [paper Methods; quote: "harvested on day 4"] ;; '
    "unknown_axis = ignored"
)
REPLICATES = 'three replicates [paper Methods; quote: "three biological replicates"]'


def crosswalk_file(tmp_path, rows):
    """Write a crosswalk with the real header; rows are (utex, relationship, source, ambiguity)."""
    path = tmp_path / "crosswalk.tsv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(CROSSWALK_HEADER)
        for utex, relationship, source, ambiguity in rows:
            writer.writerow([utex, relationship, "x", source, "seq", 1, 2, "+", ambiguity, "s", "", "m"])
    return path


def conditions_file(tmp_path):
    path = tmp_path / "conditions.tsv"
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t")
        writer.writerow(["strain", "conditions", "replicates"])
        writer.writerow(["other", "CONDITION SET: other", "n/a"])
        writer.writerow(["PCC 7942", CONDITIONS, REPLICATES])
    return path


def csv_bytes(header, rows, prefix=""):
    buffer = io.StringIO()
    buffer.write(prefix)
    writer = csv.writer(buffer)
    writer.writerow(header)
    writer.writerows(rows)
    return buffer.getvalue().encode("utf-8")


def spec_for(tmp_path, data, reader, layers=None, **overrides):
    target = tmp_path / "deposit.csv"
    target.write_bytes(data)
    spec = {
        "studyId": "GSE1", "dataType": "transcriptomics", "platform": "RNA-seq", "strain": "PCC 7942",
        "basis": "transferred", "organism": "Synechococcus elongatus PCC 7942",
        "assay": "RNA-seq transcript abundance", "licence": "CC BY 4.0", "provenanceDoc": "doc.md",
        "archiveUrl": "https://example.org/GSE1",
        "citationId": "test-2026",
        "citation": {"text": "Test 2026", "url": "https://doi.org/10.1/x", "pmid": "1"},
        "file": {"name": target.name, "url": "https://example.org/deposit.csv", "sha256": ingest.sha256_of(data)},
        "reader": reader, "units": "mean CPM", "normalization": "cpm", "conditionTableRow": 2,
        "replicates": {"count": 3, "text": "three replicates"},
        "conditions": {
            "temperature": {"status": "reported", "lo": 37, "hi": 37, "unit": "°C", "text": "37 °C"},
            "lightIntensity": {"status": "not reported", "lo": None, "hi": None,
                               "unit": "µmol photons m⁻² s⁻¹", "text": "not reported"},
            "lightRegime": {"status": "not reported", "kind": None, "photoperiod": None,
                            "spectrumClass": None, "entrained": False, "text": "not reported"},
            "co2": {"status": "reported", "lo": 0.04, "hi": 0.04, "unit": "%", "text": "ambient air"},
            "medium": {"status": "reported", "base": "BG-11", "modified": False, "conditioned": False,
                       "nitrogenAltered": False, "text": "BG-11"},
            "format": {"status": "reported", "value": "planktonic liquid", "text": "liquid"},
            "phase": {"status": "reported", "label": None, "od": None, "odNm": None, "text": "day 4"},
        },
        "caveat": "Measured in PCC 7942.",
        "layers": layers or [{
            "id": "GSE1_control", "metricKey": "exprGse1Control", "label": "Expression GSE1 control",
            "conditionSet": "control, mean of three", "samples": "GSM1-GSM3",
            "columns": ["a", "b", "c"], "treatments": [], "group": "standard",
        }],
    }
    spec.update(overrides)
    return spec


def run(tmp_path, spec, crosswalk_rows, conditions=None):
    manifest = tmp_path / "sources.json"
    out = tmp_path / "out"
    out.mkdir(exist_ok=True)
    written = ingest.ingest(
        spec, manifest_path=manifest, crosswalk_path=crosswalk_file(tmp_path, crosswalk_rows),
        interim=tmp_path, conditions_table=conditions, out_dir=out,
    )
    return written, manifest, out


def test_fetch_reads_a_cached_file_and_refuses_a_checksum_mismatch(tmp_path):
    target = tmp_path / "cached.bin"
    target.write_bytes(b"hello")
    assert ingest.fetch("https://example.org/x", hashlib.sha256(b"hello").hexdigest(), target) == b"hello"
    with pytest.raises(ValueError, match="checksum mismatch"):
        ingest.fetch("https://example.org/x", "0" * 64, target)
    assert target.read_bytes() == b"hello", "a mismatch never rewrites a staged file"


def fake_urlopen_factory(payload, calls):
    """A stand-in for urlopen that records the addresses it was asked for."""
    class Response(io.BytesIO):
        def __enter__(self):
            return self

        def __exit__(self, *exc):
            return False

    def fake_urlopen(url, timeout):
        calls.append(url)
        return Response(payload)

    return fake_urlopen


def test_fetch_never_requests_a_manually_staged_input(tmp_path, monkeypatch):
    """The Fitness Browser answers a client with a bot check, not the table.

    A missing manual input must stop with the path to stage, because requesting
    the organism page would cache its HTML challenge as the pinned dataset.
    """
    calls = []
    monkeypatch.setattr(ingest.urllib.request, "urlopen",
                        fake_urlopen_factory(b"<html>bot check</html>", calls))
    target = tmp_path / "interim" / "fit_organism_SynE.tsv"
    digest = hashlib.sha256(b"real table").hexdigest()
    with pytest.raises(ValueError, match="manually staged input"):
        ingest.fetch("https://fit.genomics.lbl.gov/x", digest, target, manual=True)
    assert calls == [], "no request is made for a manual input"
    assert not target.exists() and not target.parent.exists(), "nothing is cached"

    # Staged by the owner, it is read without a request.
    target.parent.mkdir(parents=True)
    target.write_bytes(b"real table")
    assert ingest.fetch("https://fit.genomics.lbl.gov/x", digest, target, manual=True) == b"real table"
    assert calls == []


def test_fetch_caches_a_download_only_after_its_checksum_matches(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(ingest.urllib.request, "urlopen",
                        fake_urlopen_factory(b"<html>bot check</html>", calls))
    target = tmp_path / "interim" / "file.bin"
    with pytest.raises(ValueError, match="the download was not cached"):
        ingest.fetch("https://example.org/file", "0" * 64, target)
    assert calls == ["https://example.org/file"]
    assert not target.exists(), "a failed download must not become the pinned input"


def test_fetch_downloads_once_into_the_interim_directory(tmp_path, monkeypatch):
    calls = []
    monkeypatch.setattr(ingest.urllib.request, "urlopen",
                        fake_urlopen_factory(b"payload", calls))
    target = tmp_path / "nested" / "file.bin"
    digest = hashlib.sha256(b"payload").hexdigest()
    assert ingest.fetch("https://example.org/file", digest, target) == b"payload"
    assert ingest.fetch("https://example.org/file", digest, target) == b"payload"
    assert calls == ["https://example.org/file"]
    assert target.read_bytes() == b"payload"


def test_read_table_finds_the_header_and_moves_the_identifier_first():
    data = csv_bytes(["sample", "locus_tag", "note"], [["1", "g1", "x"], ["2", "", "blank id dropped"], ["3", "g3", ""]],
                     prefix="\ufeff# a title row\n")
    header, rows = ingest.read_table(data, {"format": "csv", "idColumn": "locus_tag"})
    assert header == ["locus_tag", "sample", "note"]
    assert rows == [["g1", "1", "x"], ["g3", "3", ""]]


def test_read_table_handles_gzip_tsv_and_workbooks(tmp_path):
    tsv = gzip.compress(b"id\tv\ng1\t4\n")
    assert ingest.read_table(tsv, {"format": "tsv", "idColumn": "id"}) == (["id", "v"], [["g1", "4"]])

    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "counts"
    sheet.append(["title only"])
    sheet.append(["v", "id"])
    sheet.append([2.5, "g1"])
    sheet.append([None, "g2"])
    buffer = io.BytesIO()
    book.save(buffer)
    header, rows = ingest.read_table(buffer.getvalue(), {"format": "xlsx", "sheet": "counts", "idColumn": "id"})
    assert header == ["id", "v"]
    assert rows == [["g1", "2.5"], ["g2", ""]]


def test_read_table_joins_replicate_sheets_and_prefixes_block_titles():
    book = openpyxl.Workbook()
    first = book.active
    first.title = "Replicate 1"
    first.append(["", "", "", "", "Pulse block", ""])          # row 1: a title only over the pulse columns
    first.append([])                                          # row 2
    first.append(["", "Day block", "", "", "0", "15"])        # row 3: day title; minute offsets under the pulse
    first.append(["Gene ID", "0.5", "2", "", "8", "8.25"])    # header row
    first.append(["g1", 1, 2, "", 5, 6])
    first.append(["g2", 3, 4, "", 7, 8])
    second = book.create_sheet("Replicate 2")
    second.append(["", "", "", "", "Pulse block", ""])
    second.append([])
    second.append(["", "Day block", "", "", "0", "15"])
    second.append(["Gene ID", "0.5", "2", "", "8", "8.25"])
    second.append(["g1", 10, 20, "", 50, 60])
    buffer = io.BytesIO()
    book.save(buffer)
    reader = {"format": "xlsx", "sheets": ["Replicate 1", "Replicate 2"], "idColumn": "Gene ID", "blockTitleRows": [1, 3]}
    header, rows = ingest.read_table(buffer.getvalue(), reader)
    assert header == [
        "Gene ID",
        "Day block :: 0.5 :: Replicate 1", "Day block :: 2 :: Replicate 1", "",
        "Pulse block :: 8 :: Replicate 1", "Pulse block :: 8.25 :: Replicate 1",
        "Day block :: 0.5 :: Replicate 2", "Day block :: 2 :: Replicate 2", "",
        "Pulse block :: 8 :: Replicate 2", "Pulse block :: 8.25 :: Replicate 2",
    ]
    assert rows[0] == ["g1", "1", "2", "", "5", "6", "10", "20", "", "50", "60"]
    assert rows[1] == ["g2", "3", "4", "", "7", "8", "", "", "", "", ""], "a gene one sheet lacks keeps blank cells there"
    # One title row still works through the older key, and a single sheet carries no suffix.
    header, _ = ingest.read_table(buffer.getvalue(), {"format": "xlsx", "sheet": "Replicate 1", "idColumn": "Gene ID", "blockTitleRow": 3})
    assert header[1] == "Day block :: 0.5"
    assert header[4] == "0 :: 8", "without the first title row the minute offset is the nearest title"


def test_read_table_accepts_an_unnamed_identifier_column():
    data = csv_bytes(["", "Se_ax_90_1", "Se_ax_90_2"], [["SYNPCC7942_RS00005", "16290", "17761"], ["", "1", "2"]])
    header, rows = ingest.read_table(data, {"format": "csv", "idColumn": ""})
    assert header == ["", "Se_ax_90_1", "Se_ax_90_2"]
    assert rows == [["SYNPCC7942_RS00005", "16290", "17761"]], "a row with no identifier is dropped"


def test_select_identifiers_keeps_pattern_matches_and_names_them_by_the_group():
    rows = [["gene-SYNPCC7942_RS13335", "1"], ["rna-SYNPCC7942_RS00445", "2"], ["MSTRG.17.1", "3"], [" gene-SYNPCC7942_RS13640 ", "4"]]
    assert ingest.select_identifiers(rows, None) == (rows, 0)
    kept, dropped = ingest.select_identifiers(rows, r"gene-(SYNPCC7942_RS\d+)")
    assert kept == [["SYNPCC7942_RS13335", "1"], ["SYNPCC7942_RS13640", "4"]]
    assert dropped == 2
    # Without a group the whole match is the identifier, and a repeated
    # non-gene label is dropped before it could be read as a duplicate.
    kept, dropped = ingest.select_identifiers(
        [["Synpcc7942_0001", "1"], ["predicted RNA", "2"], ["predicted RNA", "3"]], r"Synpcc7942_\d{4}")
    assert kept == [["Synpcc7942_0001", "1"]]
    assert dropped == 2


def test_ingest_counts_identifiers_outside_the_pattern_as_unmapped(tmp_path):
    data = csv_bytes(["transcript_id", "a", "b", "c"], [
        ["gene-S1", "10", "20", "30"], ["rna-S1", "1", "1", "1"], ["MSTRG.1.1", "1", "1", "1"]])
    spec = spec_for(tmp_path, data, {"format": "csv", "idColumn": "transcript_id", "idKind": "pcc7942_rs", "idPattern": r"gene-(S\d+)"})
    written, _, out = run(tmp_path, spec, [("U1", "pcc7942_ortholog", "S1", "")], conditions_file(tmp_path))
    assert (out / "GSE1_control.tsv").read_text(encoding="utf-8").splitlines()[1] == "U1\t1000000.0000\tS1"
    assert written[0]["ingest"]["mappedGenes"] == 1
    assert written[0]["ingest"]["unmappedIdentifiers"] == 2, "the rna- and MSTRG rows"


def test_fitness_browser_reader_keeps_exact_suffix_and_plasmid_joins(tmp_path):
    """The production reader must accept every namespace form in its crosswalk."""
    production = json.loads((ROOT / "data/expression/ingest/FITNESS_BROWSER_SynE.json").read_text())
    data = csv_bytes(["locusId", "a", "b", "c"], [
        ["Synpcc7942_0001", "1", "2", "3"],
        ["Synpcc7942_1912a", "-3", "-2", "-1"],
        ["Synpcc7942_B2633", "0", "0", "0"],
        ["Synpcc7942_B9999", "9", "9", "9"],
        ["Synpcc7942_B9998", "9", "9", "9"],
    ])
    reader = {**production["reader"], "format": "csv"}
    spec = spec_for(tmp_path, data, reader, dataType="fitness", platform="RB-TnSeq", signed=True,
                    normalization="as-deposited")
    crosswalk = [
        ("U1", "pcc7942_old_locus_tag", "Synpcc7942_0001", ""),
        ("U2", "pcc7942_old_locus_tag", "Synpcc7942_1912a", ""),
        ("U3", "pcc7942_old_locus_tag", "Synpcc7942_B2633", ""),
        ("U4", "pcc7942_old_locus_tag", "Synpcc7942_B9998", "ambiguous"),
    ]
    written, _, out = run(tmp_path, spec, crosswalk, conditions_file(tmp_path))
    # A fitness spec heads its values with the fitness column, and the negative
    # and zero scores pass through untouched.
    assert (out / "GSE1_control.tsv").read_text().splitlines() == [
        "locus_tag\tfitness\tsource_gene_id",
        "U1\t2.0000\tSynpcc7942_0001",
        "U2\t-2.0000\tSynpcc7942_1912a",
        "U3\t0.0000\tSynpcc7942_B2633",
    ]
    assert written[0]["ingest"]["mappedGenes"] == 3
    assert written[0]["ingest"]["unmappedIdentifiers"] == 2


def test_column_values_refuses_a_header_that_names_several_columns():
    with pytest.raises(ValueError, match="names 2 columns"):
        ingest.column_values(["id", "a", "a"], [["g1", "1", "2"]], "a")


def test_a_layer_may_carry_its_own_condition_record_and_table_row(tmp_path):
    data = csv_bytes(["locus_tag", "a", "b"], [["S1", "1", "3"]])
    layers = [
        {"id": "GSE1_a", "metricKey": "exprGse1A", "label": "a", "conditionSet": "a", "samples": "GSM1",
         "columns": ["a"], "treatments": [], "group": "standard"},
        {"id": "GSE1_b", "metricKey": "exprGse1B", "label": "b", "conditionSet": "b", "samples": "GSM2",
         "columns": ["b"], "treatments": ["high light"], "group": "elevated", "conditionTableRow": 3,
         "conditions": None},
    ]
    spec = spec_for(tmp_path, data, {"format": "csv", "idColumn": "locus_tag", "idKind": "pcc7942_old"}, layers)
    own = copy.deepcopy(spec["conditions"])
    own["lightIntensity"] = {"status": "reported", "lo": 500, "hi": 500, "unit": "µmol photons m⁻² s⁻¹",
                             "text": "500", "where": "paper", "quote": "500 µmol"}
    spec["layers"][1]["conditions"] = own
    # A third table row, with its own quotes, for the layer that names it.
    table = conditions_file(tmp_path)
    with table.open("a", encoding="utf-8", newline="") as handle:
        csv.writer(handle, delimiter="\t").writerow(["PCC 7942", CONDITIONS.replace("37°C", "42°C").replace("Grown at 37℃", "Grown at 42℃"), REPLICATES])
    written, _, _ = run(tmp_path, spec, [("U1", "pcc7942_old_locus_tag", "S1", "")], table)
    assert written[0]["record"]["conditionTableRow"] == 2
    assert written[0]["record"]["conditions"]["lightIntensity"]["status"] == "not reported"
    assert written[0]["record"]["conditions"]["temperature"]["quote"] == "Grown at 37℃"
    assert written[1]["record"]["conditionTableRow"] == 3
    assert written[1]["record"]["conditions"]["lightIntensity"]["lo"] == 500
    assert written[1]["record"]["conditions"]["temperature"]["quote"] == "Grown at 42℃", "the layer's own row supplies its quotes"


def test_a_layer_may_describe_its_own_replicates(tmp_path):
    """One deposit can mix culture formats, so the replicate description is per layer.

    The layer's text is used as written and still takes the condition table's
    location, and a layer that says nothing keeps the study-wide description.
    """
    spec = spec_for(tmp_path, b"", {"format": "csv", "idColumn": "id", "idKind": "pcc7942_old"})
    quotes, replicate_quote = ingest.condition_quotes(conditions_file(tmp_path), 2)
    shared = ingest.build_record(spec, spec["layers"][0], quotes, replicate_quote)
    assert shared["replicates"] == {
        "count": 3, "text": "three replicates", "where": "paper Methods"}

    own = {**spec["layers"][0],
           "replicates": {"count": None, "text": "vessel count not stated"}}
    record = ingest.build_record(spec, own, quotes, replicate_quote)
    assert record["replicates"] == {
        "count": None, "text": "vessel count not stated", "where": "paper Methods"}
    assert spec["replicates"] == {"count": 3, "text": "three replicates"}, "the spec is not mutated"


def test_join_tables_joins_parts_on_the_identifier_and_labels_their_columns():
    parts = [
        ("A", ["id", "ctrl", ""], [["g1", "1", "x"], ["g2", "2", "x"]]),
        ("B", ["id", "ctrl"], [["g1", "10"], ["g3", "30"]]),
    ]
    header, rows = ingest.join_tables(parts)
    assert header == ["id", "ctrl :: A", "", "ctrl :: B"], "an unnamed column stays unnamed"
    assert rows == [
        ["g1", "1", "x", "10"],
        ["g2", "2", "x", ""],      # part B does not list g2
        ["g3", "", "", "30"],      # part A does not list g3
    ]


def test_ingest_joins_several_deposited_files_into_one_layer(tmp_path):
    """A deposit that splits its replicates over two files averages across both."""
    first = csv_bytes(["locus_tag", "ctrl"], [["S1", "10"], ["S2", "90"]])
    second = csv_bytes(["locus_tag", "ctrl"], [["S1", "30"], ["S2", "70"]])
    (tmp_path / "part_A.csv").write_bytes(first)
    (tmp_path / "part_B.csv").write_bytes(second)
    layers = [{"id": "GSE1_control", "metricKey": "exprGse1Control", "label": "control",
               "conditionSet": "control, mean of two replicate files", "samples": "GSM1-2",
               "columns": ["ctrl :: part_A", "ctrl :: part_B"], "treatments": [], "group": "standard"}]
    spec = spec_for(tmp_path, first, {"format": "csv", "idColumn": "locus_tag", "idKind": "pcc7942_old"}, layers,
                    normalization="as-deposited")
    del spec["file"]
    spec["files"] = [
        {"name": "part_A.csv", "url": "https://example.org/a.csv", "sha256": ingest.sha256_of(first)},
        {"name": "part_B.csv", "url": "https://example.org/b.csv", "sha256": ingest.sha256_of(second)},
    ]
    written, _, out = run(tmp_path, spec, [("U1", "pcc7942_old_locus_tag", "S1", ""),
                                           ("U2", "pcc7942_old_locus_tag", "S2", "")], conditions_file(tmp_path))
    assert (out / "GSE1_control.tsv").read_text(encoding="utf-8").splitlines() == [
        "locus_tag\tabundance\tsource_gene_id", "U1\t20.0000\tS1", "U2\t80.0000\tS2"]
    info = written[0]["ingest"]
    assert info["sourceFile"] == "part_A.csv; part_B.csv"
    assert info["sourceSha256"] == f"{ingest.sha256_of(first)}; {ingest.sha256_of(second)}"


def test_a_layer_may_name_its_own_strain(tmp_path):
    """One deposit may hold several genotypes; the record names the layer's own."""
    data = csv_bytes(["locus_tag", "a", "b"], [["S1", "1", "3"]])
    layers = [
        {"id": "GSE1_wt", "metricKey": "exprGse1Wt", "label": "wt", "conditionSet": "wild type", "samples": "GSM1",
         "columns": ["a"], "treatments": [], "group": "standard"},
        {"id": "GSE1_mutant", "metricKey": "exprGse1Mutant", "label": "mutant", "conditionSet": "mutant", "samples": "GSM2",
         "columns": ["b"], "treatments": [], "group": "engineered", "strain": "PCC 7942 OX-D53E"},
    ]
    spec = spec_for(tmp_path, data, {"format": "csv", "idColumn": "locus_tag", "idKind": "pcc7942_old"}, layers)
    written, _, _ = run(tmp_path, spec, [("U1", "pcc7942_old_locus_tag", "S1", "")], conditions_file(tmp_path))
    assert written[0]["record"]["strain"] == "PCC 7942", "a layer that names no strain keeps the spec's"
    assert written[1]["record"]["strain"] == "PCC 7942 OX-D53E"


def test_read_table_opens_a_member_of_a_zip_deposit():
    import zipfile
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("report/other.tsv", "id\tv\nx\t1\n")
        archive.writestr("report/prot_quant.tsv", "PG.ProteinGroups\t[1] run.PG.Quantity\nO05161\t12.5\n")
    header, rows = ingest.read_table(buffer.getvalue(), {"format": "tsv", "idColumn": "PG.ProteinGroups", "zipMember": "report/prot_quant.tsv"})
    assert header == ["PG.ProteinGroups", "[1] run.PG.Quantity"]
    assert rows == [["O05161", "12.5"]]


def test_uniprot_accessions_map_through_ordered_locus_names_one_to_one(tmp_path):
    table = tmp_path / "uniprot.tsv"
    table.write_text(
        "Entry\tReviewed\tGene Names (ordered locus)\tProtein names\tRefSeq\n"
        "O05161\treviewed\tSynpcc7942_0001\tA\tWP_1;\n"
        "O05347\tunreviewed\tSynpcc7942_0002 Synpcc7942_0003\tB\t\n"      # two loci: dropped
        "O06865\tunreviewed\t\tC\t\n"                                     # no locus
        "O06866\tunreviewed\tSynpcc7942_0004\tD\t\n"
        "O06867\tunreviewed\tSynpcc7942_0004\tE\t\n"                      # locus named twice: both dropped
        "O07345\tunreviewed\tSYNPCC7942_RS00005\tF\t\n",                   # not an old locus tag
        encoding="utf-8",
    )
    names = ingest.load_uniprot_orf_names(table)
    assert names == {"O05161": "Synpcc7942_0001"}
    values = {"O05161": 1.0, "O05347": 2.0, "Q31S88;Q31S89": 3.0, "O06866": 4.0, "P99999": 5.0}
    mapped, unmapped = ingest.through_uniprot(values, names)
    assert mapped == {"Synpcc7942_0001": 1.0}
    assert unmapped == 4


def test_ingest_routes_a_protein_deposit_through_uniprot(tmp_path):
    import zipfile
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("prot_quant.tsv", "PG.ProteinGroups\t[1] a.PG.Quantity\t[2] b.PG.Quantity\nO05161\t10\t30\nQ1;Q2\t5\t5\nO06866\tNaN\t8\n")
    data = buffer.getvalue()
    uniprot = tmp_path / "uniprot.tsv"
    uniprot.write_text("Entry\tReviewed\tGene Names (ordered locus)\tProtein names\tRefSeq\nO05161\treviewed\tS1\tA\t\nO06866\treviewed\tS2\tB\t\n".replace("S1", "Synpcc7942_0001").replace("S2", "Synpcc7942_0002"), encoding="utf-8")
    layers = [{"id": "PXD1_all", "metricKey": "protPxd1", "label": "protein", "conditionSet": "all runs", "samples": "runs 1-2",
               "columns": ["[1] a.PG.Quantity", "[2] b.PG.Quantity"], "treatments": [], "group": "standard"}]
    spec = spec_for(tmp_path, data, {"format": "tsv", "idColumn": "PG.ProteinGroups", "idKind": "uniprot_pcc7942", "zipMember": "prot_quant.tsv"}, layers,
                    dataType="proteomics", platform="LC-MS/MS", normalization="as-deposited")
    spec["conditions"]["temperature"]["where"] = "GEO growth protocol"
    manifest = tmp_path / "sources.json"; out = tmp_path / "out"; out.mkdir()
    written = ingest.ingest(spec, manifest_path=manifest, crosswalk_path=crosswalk_file(tmp_path, [
        ("U1", "pcc7942_old_locus_tag", "Synpcc7942_0001", ""), ("U2", "pcc7942_old_locus_tag", "Synpcc7942_0002", "")]),
        interim=tmp_path, conditions_table=conditions_file(tmp_path), out_dir=out, uniprot_table=uniprot)
    table = (out / "PXD1_all.tsv").read_text(encoding="utf-8").splitlines()
    assert table == ["locus_tag\tabundance\tsource_gene_id", "U1\t20.0000\tSynpcc7942_0001"], "a NaN sample drops the gene from the mean; the ;-group is ambiguous"
    info = written[0]["ingest"]
    assert info["mappedGenes"] == 1 and info["unmappedIdentifiers"] == 1
    assert info["mappingRoute"].startswith("UniProt accession to ordered locus name in uniprot.tsv, then pcc7942_old_locus_tag")


def test_read_table_rejects_unknown_formats_and_missing_headers():
    with pytest.raises(ValueError, match="unknown table format"):
        ingest.read_table(b"x", {"format": "parquet", "idColumn": "id"})
    with pytest.raises(ValueError, match="no header row names 'id'"):
        ingest.read_table(b"a,b\n1,2\n", {"format": "csv", "idColumn": "id"})


def test_column_values_skips_blanks_and_rejects_bad_cells():
    header = ["id", "a", "b"]
    rows = [["g1", "1", "NA"], ["g2", "", "2"], ["g3 ", "3"]]
    assert ingest.column_values(header, rows, "a") == {"g1": 1.0, "g3": 3.0}
    assert ingest.column_values(header, rows, "b") == {"g2": 2.0}
    with pytest.raises(ValueError, match="not in the table"):
        ingest.column_values(header, rows, "zzz")
    with pytest.raises(ValueError, match="invalid value"):
        ingest.column_values(header, [["g1", "-1"]], "a")
    with pytest.raises(ValueError, match="invalid value"):
        ingest.column_values(header, [["g1", "inf"]], "a")
    with pytest.raises(ValueError, match="appears twice"):
        ingest.column_values(header, [["g1", "1"], ["g1", "2"]], "a")


def test_a_signed_column_keeps_negative_values():
    header = ["id", "f"]
    rows = [["g1", "-1.25"], ["g2", "0.5"]]
    assert ingest.column_values(header, rows, "f", signed=True) == {"g1": -1.25, "g2": 0.5}
    with pytest.raises(ValueError, match="invalid value"):
        ingest.column_values(header, rows, "f")
    assert ingest.layer_means(header, rows, ["f"], "as-deposited", signed=True) == {"g1": -1.25, "g2": 0.5}


def test_normalise_scales_to_counts_per_million_or_keeps_values():
    values = {"g1": 1.0, "g2": 3.0}
    assert ingest.normalise(values, "as-deposited") is values
    assert ingest.normalise(values, "cpm") == {"g1": 250_000.0, "g2": 750_000.0}
    with pytest.raises(ValueError, match="no counts"):
        ingest.normalise({"g1": 0.0}, "cpm")
    with pytest.raises(ValueError, match="unknown normalization"):
        ingest.normalise(values, "tpm")


def test_layer_means_average_over_shared_identifiers_only():
    header = ["id", "a", "b"]
    rows = [["g1", "1", "3"], ["g2", "3", ""], ["g3", "", "1"]]
    assert ingest.layer_means(header, rows, ["a", "b"], "as-deposited") == {"g1": 2.0}
    assert ingest.layer_means(header, rows, [], "cpm") == {}


def test_load_crosswalk_keeps_one_to_one_rows_only(tmp_path):
    path = crosswalk_file(tmp_path, [
        ("U1", "pcc7942_old_locus_tag", "S1", ""),
        ("U2", "pcc7942_old_locus_tag", "S2", "ambiguous"),
        ("U3", "pcc7942_old_locus_tag", "S3", ""),
        ("U4", "pcc7942_old_locus_tag", "S3", ""),
        ("U5", "pcc7942_old_locus_tag", "S5", ""),
        ("U5", "pcc7942_old_locus_tag", "S6", ""),
        ("U7", "pcc7942_ortholog", "S7", ""),
    ])
    assert ingest.load_crosswalk(path, "pcc7942_old_locus_tag") == {"S1": "U1"}
    assert ingest.load_crosswalk(path, "pcc7942_ortholog") == {"S7": "U7"}


def test_map_to_utex_counts_unmapped_identifiers():
    mapped, unmapped = ingest.map_to_utex({"S1": 1.0, "S9": 2.0}, {"S1": "U1"})
    assert mapped == {"U1": (1.0, "S1")}
    assert unmapped == 1


def test_write_table_uses_the_pipeline_format_and_returns_its_digest(tmp_path):
    path = tmp_path / "layer.tsv"
    digest = ingest.write_table(path, {"U2": (2.0, "S2"), "U1": (1.23456, "S1")}, "transcriptomics")
    content = path.read_text(encoding="utf-8")
    assert content == "locus_tag\tabundance\tsource_gene_id\nU1\t1.2346\tS1\nU2\t2.0000\tS2\n"
    assert digest == hashlib.sha256(content.encode("utf-8")).hexdigest()


def test_write_table_names_a_fitness_column_and_keeps_signs_and_zero(tmp_path):
    """The downloaded field says what it holds; a negative or zero score is written as it is."""
    path = tmp_path / "fit.tsv"
    ingest.write_table(path, {"U1": (-1.25, "S1"), "U2": (0.0, "S2"), "U3": (2.5, "S3")}, "fitness")
    assert path.read_text(encoding="utf-8").splitlines() == [
        "locus_tag\tfitness\tsource_gene_id",
        "U1\t-1.2500\tS1", "U2\t0.0000\tS2", "U3\t2.5000\tS3",
    ]
    with pytest.raises(ValueError, match="no value column is defined"):
        ingest.write_table(tmp_path / "x.tsv", {"U1": (1.0, "S1")}, "metabolomics")


def test_condition_quotes_reads_axis_and_replicate_citations(tmp_path):
    path = conditions_file(tmp_path)
    quotes, replicates = ingest.condition_quotes(path, 2)
    assert quotes == {
        "temperature": {"quote": "Grown at 37℃", "where": "GEO growth protocol"},
        "lightIntensity": {"quote": "", "where": ""},
        "co2": {"quote": "ambient air", "where": "paper Methods"},
        "medium": {"quote": "in BG11 liquid medium", "where": "GEO growth protocol"},
        "format": {"quote": "liquid medium", "where": "GEO growth protocol"},
        "phase": {"quote": "harvested on day 4", "where": "paper Methods"},
    }
    assert replicates == {"quote": "three biological replicates", "where": "paper Methods"}
    assert ingest.condition_quotes(path, 1) == ({}, {"quote": "", "where": ""})
    assert ingest.condition_quotes(None, 2) == ({}, {})
    assert ingest.condition_quotes(path, None) == ({}, {})
    assert ingest.condition_quotes(tmp_path / "missing.tsv", 2) == ({}, {})


def test_build_record_attaches_quotes_without_overriding_the_spec(tmp_path):
    spec = spec_for(tmp_path, b"", {"format": "csv", "idColumn": "id", "idKind": "pcc7942_old"})
    spec["conditions"]["co2"]["quote"] = "spec quote"
    quotes, _ = ingest.condition_quotes(conditions_file(tmp_path), 2)
    record = ingest.build_record(spec, spec["layers"][0], quotes, {"quote": "x", "where": "Methods"})
    assert record["conditions"]["temperature"]["quote"] == "Grown at 37℃"
    assert record["conditions"]["temperature"]["where"] == "GEO growth protocol"
    assert record["conditions"]["co2"]["quote"] == "spec quote"
    assert record["conditions"]["co2"]["where"] == "paper Methods"
    assert record["conditions"]["lightIntensity"] == {**spec["conditions"]["lightIntensity"], "quote": "", "where": ""}
    assert record["replicates"] == {"count": 3, "text": "three replicates", "where": "Methods"}
    assert record["conditionTableRow"] == 2
    assert record["treatments"] == []
    with pytest.raises(ValueError):
        ingest.build_record(spec, {**spec["layers"][0], "group": "misc"}, {}, {})


def test_ingest_writes_layers_and_updates_the_manifest_in_place(tmp_path):
    data = csv_bytes(["locus_tag", "a", "b", "c", "d"], [
        ["S1", "10", "20", "30", "1"],
        ["S2", "90", "80", "70", "1"],
        ["S3", "0", "0", "0", "1"],   # drops out of the crosswalk below
    ])
    layers = [
        {"id": "GSE1_control", "metricKey": "exprGse1Control", "label": "control", "conditionSet": "control",
         "samples": "GSM1-3", "columns": ["a", "b", "c"], "treatments": [], "group": "standard"},
        {"id": "GSE1_salt", "metricKey": "exprGse1Salt", "label": "salt", "conditionSet": "salt",
         "samples": "GSM4", "columns": ["d"], "treatments": ["salt"], "group": "stress", "caveat": "salt caveat"},
    ]
    spec = spec_for(tmp_path, data, {"format": "csv", "idColumn": "locus_tag", "idKind": "pcc7942_old"}, layers)
    rows = [("U1", "pcc7942_old_locus_tag", "S1", ""), ("U2", "pcc7942_old_locus_tag", "S2", "")]
    written, manifest, out = run(tmp_path, spec, rows, conditions_file(tmp_path))

    assert [entry["id"] for entry in written] == ["GSE1_control", "GSE1_salt"]
    control = (out / "GSE1_control.tsv").read_text(encoding="utf-8").splitlines()
    assert control == ["locus_tag\tabundance\tsource_gene_id", "U1\t200000.0000\tS1", "U2\t800000.0000\tS2"]
    salt = (out / "GSE1_salt.tsv").read_text(encoding="utf-8").splitlines()
    assert salt == ["locus_tag\tabundance\tsource_gene_id", "U1\t333333.3333\tS1", "U2\t333333.3333\tS2"]

    entry = written[0]
    assert entry["file"] == "GSE1_control.tsv"
    assert entry["isTargetOrganism"] is False
    assert entry["citationId"] == "test-2026"
    assert entry["payload"] == "expression_layers.json"
    assert entry["condition"] == "control"
    assert entry["caveat"] == "Measured in PCC 7942."
    assert written[1]["caveat"] == "salt caveat"
    assert entry["sha256"] == ingest.sha256_of((out / "GSE1_control.tsv").read_bytes())
    assert entry["ingest"] == {
        "sourceFile": "deposit.csv", "sourceSha256": spec["file"]["sha256"], "sourceUrl": spec["file"]["url"],
        "columns": ["a", "b", "c"], "normalization": "cpm", "mappedGenes": 2, "unmappedIdentifiers": 1,
        "mappingRoute": "pcc7942_old_locus_tag in crosswalk.tsv, one-to-one rows only",
    }
    assert entry["record"]["conditions"]["temperature"]["quote"] == "Grown at 37℃"
    assert entry["record"]["replicates"]["where"] == "paper Methods"
    assert [e["id"] for e in json.loads(manifest.read_text(encoding="utf-8"))] == ["GSE1_control", "GSE1_salt"]

    # A second run replaces the existing entries instead of appending duplicates.
    spec["layers"] = [{**layers[1], "samples": "GSM4 rerun"}]
    written, manifest, out = run(tmp_path, spec, rows, conditions_file(tmp_path))
    stored = json.loads(manifest.read_text(encoding="utf-8"))
    assert [e["id"] for e in stored] == ["GSE1_control", "GSE1_salt"]
    assert stored[1]["record"]["samples"] == "GSM4 rerun"


def test_ingest_refuses_a_layer_that_maps_nothing(tmp_path):
    data = csv_bytes(["locus_tag", "a", "b", "c"], [["S1", "1", "1", "1"]])
    spec = spec_for(tmp_path, data, {"format": "csv", "idColumn": "locus_tag", "idKind": "pcc7942_rs"})
    with pytest.raises(ValueError, match="maps no gene"):
        run(tmp_path, spec, [("U1", "pcc7942_old_locus_tag", "S1", "")])


def test_main_runs_a_spec_file_and_reports_each_layer(tmp_path, capsys):
    data = csv_bytes(["locus_tag", "a", "b", "c"], [["S1", "1", "2", "3"]])
    spec = spec_for(tmp_path, data, {"format": "csv", "idColumn": "locus_tag", "idKind": "pcc7942_old"})
    spec_path = tmp_path / "spec.json"
    spec_path.write_text(json.dumps(spec), encoding="utf-8")
    manifest = tmp_path / "manifest" / "sources.json"
    manifest.parent.mkdir()
    crosswalk = crosswalk_file(tmp_path, [("U1", "pcc7942_old_locus_tag", "S1", "")])
    code = ingest.main([
        str(spec_path), "--manifest", str(manifest), "--crosswalk", str(crosswalk),
        "--interim", str(tmp_path), "--conditions", str(conditions_file(tmp_path)),
    ])
    assert code == 0
    assert capsys.readouterr().out == "GSE1_control: 1 genes mapped, 0 identifiers unmapped -> GSE1_control.tsv\n"
    assert (manifest.parent / "GSE1_control.tsv").is_file()
    assert json.loads(manifest.read_text(encoding="utf-8"))[0]["record"]["conditionTableRow"] == 2


def dtaselect_report(proteins, *, count_header="Spectrum Count"):
    """A DTASelect report: preamble, two header lines, protein lines each followed by a peptide line, summary rows."""
    lines = ["DTASelect v2.0.47", "/scratch/search", "", f"Locus\tSequence Count\t{count_header}\tSequence Coverage\tLength\tMolWt\tpI\tValidation Status\tDescriptive Name",
             "Unique\tFileName\tXCorr\tDeltCN\tConf%\tM+H+\tCalcM+H+\tTotalIntensity\tSpR\tProb Score\tIonProportion\tRedundancy\tSequence"]
    for locus, count in proteins:
        lines.append(f"{locus}\t3\t{count}\t40.0%\t163\t17288\t5.6\tU\tSome protein OS=Synechococcus elongatus")
        lines.append("*\t7.14821.14821.2\t6.357\t0.4866\t100.0\t1705.55\t1701.8766\t1947317.0\t1\t10.15\t92.3\t174\tM.SKTPLTEAVAAADSQGR.F")
        lines.append("\t7.14822.14822.2\t5.1\t0.4\t100.0\t1705.55\t1701.8766\t1947317.0\t1\t10.15\t92.3\t174\tM.SKTPLTEAVAAADSQGR.F")
    lines += ["", "Unfiltered\t5386\t322821\t325002", "Filtered\t804\t6291\t27388", "Forward FDR\t1.13\t0.32\t0.08"]
    return "\n".join(lines) + "\n"


def dtaselect_archive(members):
    import zipfile
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        for name, text in members.items():
            archive.writestr(name, text)
    return buffer.getvalue()


def test_read_dtaselect_tables_protein_counts_per_run_by_accession():
    data = dtaselect_archive({
        "search/log-1.txt": dtaselect_report([("P13530_PHCA_SYNE7", 1545), ("Reverse_Q31QQ3_Q31QQ3_SYNE7", 4),
                                              ("contaminant_sp-P04264-K2C1_HUMAN", 9), ("Q31MM4_Q31MM4_SYNE7", 2)]),
        "search/log-2.txt": dtaselect_report([("P13530_PHCA_SYNE7", 1077), ("A0A0H3K6W5_A0A0H3K6W5_SYNE7", 7)]),
    })
    header, rows = ingest.read_table(data, {"format": "dtaselect", "zipMembers": ["search/log-1.txt", "search/log-2.txt"]})
    assert header == ["Locus", "log-1", "log-2"]
    assert rows == [
        ["P13530", "1545", "1077"],
        ["contaminant_sp-P04264-K2C1_HUMAN", "9", ""],   # kept by name; fails to map later
        ["Q31MM4", "2", ""],                               # a run that did not list it leaves a blank
        ["A0A0H3K6W5", "", "7"],                           # a ten-character accession
    ], "decoys are dropped, peptide and summary lines are never proteins"
    means = ingest.layer_means(header, rows, ["log-1", "log-2"], "as-deposited")
    assert means == {"P13530": 1311.0}, "the mean covers only proteins every run identified"
    # A different count column can be named.
    header, rows = ingest.read_table(dtaselect_archive({"a.txt": dtaselect_report([("P13530_PHCA_SYNE7", 5)])}),
                                     {"format": "dtaselect", "zipMembers": ["a.txt"], "countColumn": "Sequence Count"})
    assert rows == [["P13530", "3"]]


def test_read_dtaselect_names_what_is_wrong():
    good = dtaselect_report([("P13530_PHCA_SYNE7", 5)])
    with pytest.raises(ValueError, match="two DTASelect members share the name 'a'"):
        ingest.read_dtaselect(dtaselect_archive({"x/a.txt": good, "y/a.txt": good}), ["x/a.txt", "y/a.txt"])
    with pytest.raises(ValueError, match="no DTASelect protein header names 'Spectrum Count'"):
        ingest.read_dtaselect(dtaselect_archive({"a.txt": dtaselect_report([("P13530_PHCA_SYNE7", 5)], count_header="Spectra")}), ["a.txt"])
    with pytest.raises(ValueError, match="no DTASelect protein header"):
        ingest.read_dtaselect(dtaselect_archive({"a.txt": "DTASelect v2.0.47\nnothing here\n"}), ["a.txt"])
    with pytest.raises(ValueError, match="P13530 is listed twice"):
        ingest.read_dtaselect(dtaselect_archive({"a.txt": dtaselect_report([("P13530_PHCA_SYNE7", 5), ("P13530_OTHER_SYNE7", 6)])}), ["a.txt"])
    with pytest.raises(ValueError, match="a.txt: no protein lines"):
        ingest.read_dtaselect(dtaselect_archive({"a.txt": dtaselect_report([("Reverse_P13530_PHCA_SYNE7", 5)])}), ["a.txt"])


MZTAB_PRH = ("PRH\taccession\tdescription\tnum_psms_ms_run[1]\tnum_peptides_distinct_ms_run[1]"
             "\topt_global_cv_PRIDE:0000303_Decoy_hit")


def mztab_file(proteins, *, header=MZTAB_PRH, metadata=True):
    """An mzTab identification file: metadata, the protein table, then a PSM section."""
    lines = []
    if metadata:
        lines += ["MTD\tmzTab-version\t1.0", "MTD\tmzTab-type\tIdentification",
                  "MTD\tsample[1]-description\tS. elongatus LD6.5-14.5", "COM\ta comment"]
    lines.append(header)
    for accession, psms, distinct, decoy in proteins:
        lines.append(f"PRT\t{accession}\tsome protein\t{psms}\t{distinct}\t{decoy}")
    # A PSM section follows the protein table and must never be read as protein rows.
    lines += ["PSH\tsequence\tPSM_ID\taccession", "PSM\tPEPTIDEK\t1\tO05161"]
    return "\n".join(lines).encode("utf-8")


def test_read_mztab_tables_psm_counts_by_accession():
    data = mztab_file([
        ("O05161", 62, 16, 0),
        ("Q31RN1", 4, 2, 1),          # flagged a decoy hit
        ("O06865", 73, 21, 0),
        ("O32463", "null", 0, 0),     # no count to read
        ("", 5, 1, 0),                # no accession to key on
    ])
    header, rows = ingest.read_table(data, {"format": "mztab"})
    assert header == ["accession", "num_psms"]
    assert rows == [["O05161", "62"], ["O06865", "73"]], \
        "decoys, countless rows and the metadata and PSM sections are all left out"
    means = ingest.layer_means(header, rows, ["num_psms"], "cpm")
    assert means == {"O05161": 62 / 135 * 1_000_000, "O06865": 73 / 135 * 1_000_000}, \
        "cpm scales over the proteins the run reports"
    # Gzip is detected by magic bytes, as for every other format.
    gzipped, _ = ingest.read_table(gzip.compress(data), {"format": "mztab"})
    assert gzipped == header
    # A different count column can be named.
    _, distinct = ingest.read_table(data, {"format": "mztab", "countColumn": "num_peptides_distinct_ms_run[1]"})
    assert distinct == [["O05161", "16"], ["O06865", "21"], ["O32463", "0"]], \
        "a row is dropped for the column being read, so O32463 returns once that column has a value"


def test_read_mztab_joins_one_file_per_window_under_its_label():
    parts = []
    for label, proteins in (("LD6.5-14.5", [("O05161", 62, 16, 0)]),
                            ("LD17.5-26.5", [("O05161", 90, 20, 0), ("O06865", 7, 3, 0)])):
        header, rows = ingest.read_table(mztab_file(proteins), {"format": "mztab"})
        parts.append((label, header, rows))
    header, rows = ingest.join_tables(parts)
    assert header == ["accession", "num_psms :: LD6.5-14.5", "num_psms :: LD17.5-26.5"]
    assert rows == [["O05161", "62", "90"], ["O06865", "", "7"]], \
        "a window that did not identify a protein leaves a blank, not a zero"
    # Each layer reads only its own window's column.
    assert ingest.layer_means(header, rows, ["num_psms :: LD6.5-14.5"], "cpm") == {"O05161": 1_000_000.0}


def test_read_mztab_names_what_is_wrong():
    with pytest.raises(ValueError, match="holds no PRH protein header"):
        ingest.read_mztab(b"MTD\tmzTab-version\t1.0\n")
    with pytest.raises(ValueError, match="names no accession column"):
        ingest.read_mztab(mztab_file([], header="PRH\tdescription\tnum_psms_ms_run[1]"))
    with pytest.raises(ValueError, match="names no 'num_psms_ms_run\\[1\\]' column"):
        ingest.read_mztab(mztab_file([], header="PRH\taccession\tdescription\tnum_psms"))
    with pytest.raises(ValueError, match="PRT row appears before its PRH header"):
        ingest.read_mztab(b"PRT\tO05161\tsome protein\t62\t16\t0\n")
    with pytest.raises(ValueError, match="narrower than its PRH header"):
        ingest.read_mztab(b"\n".join([MZTAB_PRH.encode(), b"PRT\tO05161\t62"]))
    with pytest.raises(ValueError, match="names 'O05161' twice"):
        ingest.read_mztab(mztab_file([("O05161", 62, 16, 0), ("O05161", 7, 2, 0)]))


MZID_NS = "http://psidev.info/psi/pi/mzIdentML/1.1"


def mzid_file(matches, *, proteins=(("DBSeq1", "sp|P06539|PHCB_SYNE7"), ("DBSeq2", "tr|Q31ML1|Q31ML1_SYNE7")),
              evidence=(("PepEv_1", "DBSeq1", "false"), ("PepEv_2", "DBSeq2", "false"),
                        ("PepEv_D", "DBSeq1", "true"))):
    """A small MS-GF+ mzIdentML 1.1 file; each match is (refs, rank, pass, qvalue)."""
    parts = [f'<?xml version="1.0" encoding="UTF-8"?>', f'<MzIdentML xmlns="{MZID_NS}" version="1.1.0">',
             "<SequenceCollection>"]
    for seq_id, accession in proteins:
        parts.append(f'<DBSequence id="{seq_id}" accession="{accession}"/>')
    for ev_id, seq_ref, decoy in evidence:
        parts.append(f'<PeptideEvidence id="{ev_id}" dBSequence_ref="{seq_ref}" isDecoy="{decoy}"/>')
    parts.append("</SequenceCollection><AnalysisData><SpectrumIdentificationList>")
    for index, (refs, rank, passes, qvalue) in enumerate(matches):
        parts.append(f'<SpectrumIdentificationResult id="SIR_{index}">')
        parts.append(f'<SpectrumIdentificationItem rank="{rank}" passThreshold="{passes}" id="SII_{index}">')
        for ref in refs:
            parts.append(f'<PeptideEvidenceRef peptideEvidence_ref="{ref}"/>')
        if qvalue is not None:
            parts.append(f'<cvParam cvRef="PSI-MS" accession="MS:1002054" name="MS-GF:QValue" value="{qvalue}"/>')
        parts.append("</SpectrumIdentificationItem></SpectrumIdentificationResult>")
    parts.append("</SpectrumIdentificationList></AnalysisData></MzIdentML>")
    return "".join(parts).encode("utf-8")


def test_read_mzidentml_counts_unique_rank_one_matches_per_protein():
    data = mzid_file([
        ((["PepEv_1"]), "1", "true", "0.0"),            # counts for P06539
        ((["PepEv_1"]), "1", "true", "0.001"),          # counts for P06539
        ((["PepEv_2"]), "1", "true", "0.0"),            # counts for Q31ML1
        ((["PepEv_1", "PepEv_2"]), "1", "true", "0.0"),  # shared peptide, counted for neither
        ((["PepEv_1"]), "2", "true", "0.0"),            # not rank 1
        ((["PepEv_1"]), "1", "false", "0.0"),           # did not pass the search threshold
        ((["PepEv_1"]), "1", "true", "0.5"),            # above the q-value ceiling
        ((["PepEv_1"]), "1", "true", None),             # no q-value reported
        ((["PepEv_D"]), "1", "true", "0.0"),            # decoy evidence only
    ])
    header, rows = ingest.read_table(data, {"format": "mzidentml"})
    assert header == ["accession", "psm_count"]
    assert rows == [["P06539", "2"], ["Q31ML1", "1"]], \
        "sp|/tr| accessions are reduced to the UniProt accession, and only unique rank-1 matches count"
    # A looser ceiling admits the 0.5 match; the shared one still counts for neither.
    _, loose = ingest.read_table(data, {"format": "mzidentml", "qValueMax": 0.6})
    assert loose == [["P06539", "3"], ["Q31ML1", "1"]]
    # Gzip is detected by magic bytes, as for every other format.
    assert ingest.read_table(gzip.compress(data), {"format": "mzidentml"})[1] == rows


def test_read_mzidentml_names_what_is_wrong():
    with pytest.raises(ValueError, match="lists no DBSequence"):
        ingest.read_mzidentml(f'<MzIdentML xmlns="{MZID_NS}" version="1.1.0"/>'.encode())
    with pytest.raises(ValueError, match="no rank-1 match at q <= 0.01"):
        ingest.read_mzidentml(mzid_file([((["PepEv_1"]), "1", "true", "0.9")]))
