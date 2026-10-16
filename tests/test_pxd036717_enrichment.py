"""Source preservation and non-imputed reconstruction; synthetic XLSX fixtures."""

from datetime import datetime
import hashlib
import io
import json
import math
from pathlib import Path
import sys
from types import SimpleNamespace
from xml.etree import ElementTree
import zipfile

from openpyxl import Workbook
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "tools"))
import pxd036717_enrichment as enrichment  # noqa: E402


def workbook(mutate=None):
    """Make literal source measurements, with optional malformed source edits."""
    book = Workbook()
    book.remove(book.active)
    for bait in enrichment.BAITS:
        sheet = book.create_sheet(f"{bait}-APEX2")
        sheet.append(enrichment.expected_headers(bait))
        sheet.append([
            "Q12345", "gene  exact ", 4, 8, 16, 16, 4, 16,
            2, 3, 4, 4, 2, 4, 2, -1, 0,
            2, 3, 4, 4, 3, 2, 12, 13, 14, 15, 16, 17, 35, 123,
            "sp|Q12345|EXACT source header",
        ])
        sheet.append([
            "Q54321", None, 0, None, -1, 0, 16, 16,
            10, 10, 10, 10, 4, 4, 0, -6, -6,
            *([None] * 15),
        ])
    if mutate:
        mutate(book)
    output = io.BytesIO()
    book.save(output)
    book.close()
    return output.getvalue()


def read(data):
    return enrichment.read_workbook(data, expected_sha256=hashlib.sha256(data).hexdigest())


def test_preservation_reconstruction_summary_and_json_round_trip():
    data = workbook()
    doc = json.loads(json.dumps(read(data), allow_nan=False))
    assert doc["source"]["sha256"] == hashlib.sha256(data).hexdigest()
    assert doc["source"]["bytes"] == len(data)
    assert not doc["releaseEligible"]
    assert [s["bait"] for s in doc["sheets"]] == list(enrichment.BAITS)
    for sheet in doc["sheets"]:
        positive, missing = sheet["rows"]
        assert positive["cells"][1]["value"] == "gene  exact "
        assert positive["cells"][31]["cell"] == "AF2"
        assert positive["cells"][31]["value"] == "sp|Q12345|EXACT source header"
        assert [c["reconstructedLog2Ratio"] for c in positive["comparisons"]] == [2, -1, 0]
        assert all(not c["flags"] for c in positive["comparisons"])
        assert positive["validReplicates"] == 3
        assert positive["meanLog2Ratio"] == pytest.approx(1 / 3)
        assert positive["comparisons"][2]["cells"] == {
            "controlIntensity": "E2", "baitIntensity": "H2",
            "publishedControlLog2": "K2", "publishedBaitLog2": "N2", "publishedRatio": "Q2",
        }
        assert missing["validReplicates"] == 0 and missing["meanLog2Ratio"] is None
        assert missing["cells"][2]["value"] == 0
        assert missing["cells"][3]["value"] is None
        assert missing["comparisons"][0]["reconstructedLog2Ratio"] is None
        assert sheet["audit"]["validByReplicate"] == {"A": 1, "B": 1, "C": 1}
        assert sheet["audit"]["completeMeans"] == 1


@pytest.mark.parametrize("raw,flag", [
    (None, "missing"), (0, "missing"), (-2, "invalid"), (True, "invalid"),
    ("12", "invalid"), (float("nan"), "invalid"), (float("inf"), "invalid"),
])
@pytest.mark.parametrize("arm", ["control", "bait"])
def test_missing_and_invalid_intensities_never_become_ratios(raw, flag, arm):
    args = [4, 16, 2, 4, 2]
    args[0 if arm == "control" else 1] = raw
    result = enrichment.reconstruct_comparison(*args)
    assert result["reconstructedLog2Ratio"] is None
    assert f"{flag}_{arm}_intensity" in result["flags"]
    assert f"published_{arm}_log2_without_positive_intensity" in result["flags"]


def test_inconsistent_source_logs_and_ratios_do_not_contaminate_reconstruction():
    result = enrichment.reconstruct_comparison(4, 16, 0, 1, 99)
    assert result["reconstructedLog2Ratio"] == 2
    assert set(result["flags"]) == {
        "published_control_log2_disagrees_with_intensity",
        "published_bait_log2_disagrees_with_intensity",
        "published_ratio_disagrees_with_published_logs",
        "published_ratio_disagrees_with_reconstruction",
    }
    unavailable = enrichment.reconstruct_comparison(4, 16, None, "#VALUE!", None)
    assert unavailable["reconstructedLog2Ratio"] == 2
    assert set(unavailable["flags"]) == {
        "unavailable_published_control_log2", "unavailable_published_bait_log2",
        "unavailable_published_ratio",
    }
    partial_log = enrichment.reconstruct_comparison(4, 16, None, 4, 2)
    assert partial_log["reconstructedLog2Ratio"] == 2
    assert partial_log["flags"] == ["unavailable_published_control_log2"]
    assert not enrichment.reconstruct_comparison(4, 16, 2, 4, 2 + 1e-8)["flags"]
    assert enrichment.reconstruct_comparison(4, 16, 2, 4, 2 + 1e-6)["flags"]
    # Subtraction of logarithms avoids overflow in a direct bait/control division.
    extreme = enrichment.reconstruct_comparison(1e-300, 1e300, None, None, None)
    assert math.isfinite(extreme["reconstructedLog2Ratio"])


def test_formulas_cached_values_errors_and_missing_caches_are_preserved():
    def mutate(book):
        book.worksheets[0]["O2"] = "=L2-I2"
        book.worksheets[0]["P2"] = "=M2-J2"
        book.worksheets[0]["Q2"] = "#VALUE!"
    original = workbook(mutate)
    output = io.BytesIO()
    # openpyxl intentionally does not calculate formulas. Set one synthetic cache
    # directly in OOXML, leaving the other absent to exercise both source states.
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(output, "w") as target:
        for name in source.namelist():
            payload = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(payload)
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                root.find(".//s:c[@r='O2']/s:v", ns).text = "2"
                payload = ElementTree.tostring(root)
            target.writestr(name, payload)
    row = read(output.getvalue())["sheets"][0]["rows"][0]
    assert row["cells"][14] == {
        "cell": "O2", "value": "=L2-I2", "type": "f", "cachedValue": 2, "cachedType": "n",
    }
    assert row["cells"][15]["value"] == "=M2-J2" and row["cells"][15]["cachedValue"] is None
    assert row["cells"][16]["type"] == "e" and row["cells"][16]["cachedValue"] == "#VALUE!"
    assert "unavailable_published_ratio" in row["comparisons"][1]["flags"]
    assert row["meanLog2Ratio"] == pytest.approx(1 / 3)


def test_duplicate_accessions_blank_rows_and_blank_extra_columns_are_retained():
    def mutate(book):
        sheet = book.worksheets[0]
        sheet["A3"] = "Q12345"
        sheet.cell(5, 33).value = None
        sheet.cell(5, 33).number_format = "0.0"
    table = read(workbook(mutate))["sheets"][0]
    assert len(table["rows"]) == 2
    assert table["blankRows"] == [4, 5]
    assert table["dimensions"] == [5, 33]
    assert table["audit"]["duplicateAccessions"] == {"Q12345": 2}


@pytest.mark.parametrize("mutate,message", [
    (lambda b: setattr(b.worksheets[0], "title", "Other"), "sheet names"),
    (lambda b: b.worksheets[0].__setitem__("A1", "UniProt identifier"), "headers"),
    (lambda b: b.worksheets[0].__setitem__("AG1", "Unexpected"), "headers"),
    (lambda b: b.worksheets[0].__setitem__("AG2", 42), "extra column"),
    (lambda b: b.worksheets[0].__setitem__("A2", None), "identifier"),
    (lambda b: b.worksheets[0].__setitem__("A2", "=A3"), "identifier"),
    (lambda b: b.worksheets[0].__setitem__("A2", "  "), "identifier"),
    (lambda b: b.worksheets[0].delete_rows(2, 2), "no data rows"),
    (lambda b: b.worksheets[0].__setitem__("B2", datetime(2026, 1, 1)), "source value type"),
])
def test_schema_changes_fail_closed(mutate, message):
    with pytest.raises(ValueError, match=message):
        read(workbook(mutate))


def test_source_hash_and_nonfinite_cells_fail_closed():
    with pytest.raises(ValueError, match="SHA-256"):
        enrichment.read_workbook(workbook())
    cell = SimpleNamespace(value=float("inf"), data_type="n")
    with pytest.raises(ValueError, match="non-finite"):
        enrichment._source_cell(cell, cell, "C2")


def test_partial_three_replicate_mean_is_missing():
    data = workbook(lambda b: b.worksheets[0].__setitem__("H2", 0))
    row = read(data)["sheets"][0]["rows"][0]
    assert row["validReplicates"] == 2
    assert row["meanLog2Ratio"] is None
    assert row["comparisons"][2]["reconstructedLog2Ratio"] is None


def test_zero_published_log_and_multiaccession_are_explicit():
    result = enrichment.reconstruct_comparison(0, 0, 0, 0, 0)
    assert result["reconstructedLog2Ratio"] is None
    assert "zero_published_control_log2_without_positive_intensity" in result["flags"]
    assert "zero_published_bait_log2_without_positive_intensity" in result["flags"]
    data = workbook(lambda b: b.worksheets[0].__setitem__("A2", "Q12345;Q99999"))
    table = read(data)["sheets"][0]
    assert table["rows"][0]["accession"] == "Q12345;Q99999"
    assert table["audit"]["multiAccessionRows"] == [2]


@pytest.mark.parametrize("control,bait,state", [
    (4, 16, "both"), (0, 16, "bait_only"), (None, 16, "bait_only"),
    (4, 0, "control_only"), (4, None, "control_only"),
    (0, 0, "neither"), (0, None, "neither"), (None, None, "neither"),
    (-1, 16, "invalid"), (4, "bad", "invalid"), (None, True, "invalid"),
])
def test_detection_is_distinct_from_ratio_missingness(control, bait, state):
    result = enrichment.reconstruct_comparison(control, bait, None, None, None)
    assert result["detectionState"] == state
    assert (result["reconstructedLog2Ratio"] is not None) == (state == "both")


def test_detection_counts_preserve_bait_only_even_when_mean_is_null():
    def mutate(book):
        for col in "CDE":
            book.worksheets[0][f"{col}2"] = 0
    table = read(workbook(mutate))["sheets"][0]
    row = table["rows"][0]
    assert row["meanLog2Ratio"] is None
    assert row["validReplicates"] == 0
    assert row["detectionCounts"] == {
        "both": 0, "bait_only": 3, "control_only": 0, "neither": 0, "invalid": 0,
    }
    assert table["audit"]["detectionCounts"] == {
        "both": 0, "bait_only": 4, "control_only": 0, "neither": 1, "invalid": 1,
    }


@pytest.mark.parametrize("count,raw,status,flags", [
    (True, None, "missing", []), (False, None, "missing", []),
    (True, 0, "zero", []), (False, 0.0, "zero", []),
    (True, 1, "positive", []), (True, 2.0, "positive", []),
    (False, 0.001, "positive", []), (False, 100, "positive", []),
    (True, 0.5, "invalid", ["non_integer"]),
    (True, -0.5, "invalid", ["negative", "non_integer"]),
    (True, -1, "invalid", ["negative"]), (False, -0.001, "invalid", ["negative"]),
    (False, 100.001, "invalid", ["above_100_percent"]),
    *[(count, raw, "invalid", ["invalid_type"])
      for count in (False, True) for raw in (False, True, "1", "", "bad")],
    *[(count, raw, "invalid", ["non_finite"])
      for count in (False, True) for raw in (float("inf"), -float("inf"), float("nan"))],
])
def test_identification_field_validation(count, raw, status, flags):
    cell = {"cell": "R2", "value": raw, "type": "n", "cachedValue": raw, "cachedType": "n"}
    result = enrichment._identification_field(cell, count=count)
    assert result == {
        "cell": "R2", "value": raw if status in ("zero", "positive") else None,
        "status": status, "flags": flags,
    }
    json.dumps(result, allow_nan=False)


def test_identification_samples_preserve_exact_cells_and_all_six_source_labels():
    doc = read(workbook())
    assert "no identification success" in doc["method"]["identificationEvidence"]["boundary"]
    for table in doc["sheets"]:
        row = table["rows"][0]
        samples = row["identificationEvidence"]
        assert len(samples) == 6
        for i, sample in enumerate(samples):
            arm = "control" if i < 3 else "bait"
            label = "control" if i < 3 else f"{table['bait']}-APEX2"
            assert sample == {
                "sheet": table["sheet"], "sample": f"{label} {'ABC'[i % 3]}",
                "arm": arm, "replicate": "ABC"[i % 3],
                "intensityCell": f"{'CDEFGH'[i]}2", "positiveLfq": True,
                "uniquePeptides": {
                    "cell": f"{'RSTUVW'[i]}2", "value": [2, 3, 4, 4, 3, 2][i],
                    "status": "positive", "flags": [],
                },
                "sequenceCoveragePercent": {
                    "cell": f"{['X', 'Y', 'Z', 'AA', 'AB', 'AC'][i]}2",
                    "value": 12 + i, "status": "positive", "flags": [],
                },
                "flags": ["spectral_counts_not_supplied"],
            }
            for key in ("uniquePeptides", "sequenceCoveragePercent"):
                source = next(cell for cell in row["cells"] if cell["cell"] == sample[key]["cell"])
                assert source["value"] == source["cachedValue"] == sample[key]["value"]
        assert table["audit"]["identificationEvidence"] == {
            "sampleCount": 12, "positiveLfqSamples": 8,
            "uniquePeptides": {"missing": 6, "zero": 0, "positive": 6, "invalid": 0},
            "sequenceCoveragePercent": {"missing": 6, "zero": 0, "positive": 6, "invalid": 0},
            "flags": {"positive_lfq_with_missing_unique_peptides": 2,
                      "positive_lfq_with_missing_sequence_coverage": 2,
                      "spectral_counts_not_supplied": 12},
        }


@pytest.mark.parametrize("index", range(6))
@pytest.mark.parametrize("raw,status", [(None, "missing"), (0, "zero"), (True, "invalid"),
                                       (-1, "invalid"), ("2", "invalid")])
def test_positive_lfq_identification_diagnostics_do_not_change_ratios(index, raw, status):
    unique_cell = f"{'RSTUVW'[index]}2"
    coverage_cell = f"{['X', 'Y', 'Z', 'AA', 'AB', 'AC'][index]}2"
    def mutate(book):
        for sheet in book:
            sheet[unique_cell] = raw
            sheet[coverage_cell] = raw
    doc = read(workbook(mutate))
    for table in doc["sheets"]:
        row = table["rows"][0]
        sample = row["identificationEvidence"][index]
        expected_flags = [f"positive_lfq_with_{status}_unique_peptides",
                          f"positive_lfq_with_{status}_sequence_coverage"]
        if status == "zero":
            expected_flags.append("below_reported_unique_peptide_minimum")
        assert sample["flags"] == expected_flags + ["spectral_counts_not_supplied"]
        for key, cell in (("uniquePeptides", unique_cell), ("sequenceCoveragePercent", coverage_cell)):
            assert sample[key]["status"] == status
            assert sample[key]["value"] == (0 if status == "zero" else None)
            assert next(c["value"] for c in row["cells"] if c["cell"] == cell) == raw
        assert [c["reconstructedLog2Ratio"] for c in row["comparisons"]] == [2, -1, 0]
        assert [c["detectionState"] for c in row["comparisons"]] == ["both"] * 3
        assert row["validReplicates"] == 3
        assert row["meanLog2Ratio"] == pytest.approx(1 / 3)
        # The second fixture row already has two positive-LFQ samples with absent evidence.
        expected_count = 3 if status == "missing" else 1
        assert table["audit"]["identificationEvidence"]["flags"][sample["flags"][0]] == expected_count


@pytest.mark.parametrize("lfq", [0, None, -1, True, "4"])
def test_nonpositive_or_invalid_lfq_does_not_raise_positive_lfq_diagnostics(lfq):
    def mutate(book):
        book.worksheets[0]["C2"] = lfq
        book.worksheets[0]["R2"] = 0
        book.worksheets[0]["X2"] = "bad"
    sample = read(workbook(mutate))["sheets"][0]["rows"][0]["identificationEvidence"][0]
    assert sample["positiveLfq"] is False
    assert sample["flags"] == ["below_reported_unique_peptide_minimum", "spectral_counts_not_supplied"]
    assert sample["uniquePeptides"]["status"] == "zero"
    assert sample["sequenceCoveragePercent"]["flags"] == ["invalid_type"]


@pytest.mark.parametrize("peptides,below", [(0, True), (1, True), (2, False), (3, False),
                                         (None, False), (-1, False), (True, False)])
def test_source_criterion_flags_never_filter_or_claim_full_identification(peptides, below):
    data = workbook(lambda b: b.worksheets[0].__setitem__("R2", peptides))
    doc = read(data)
    row = doc["sheets"][0]["rows"][0]
    sample = row["identificationEvidence"][0]
    assert ("below_reported_unique_peptide_minimum" in sample["flags"]) == below
    assert "spectral_counts_not_supplied" in sample["flags"]
    assert row["comparisons"][0]["reconstructedLog2Ratio"] == 2
    assert row["meanLog2Ratio"] == pytest.approx(1 / 3)
    assert row["cells"][17]["value"] == peptides
    assert doc["method"]["identificationEvidence"]["reportedCriteria"]["minimumUniquePeptides"] == 2
    assert not doc["releaseEligible"]


def test_identification_formula_caches_and_errors_preserve_source_and_status():
    def mutate(book):
        for col in ("R", "S", "T", "U", "V", "W", "X", "Y", "Z", "AA", "AB", "AC"):
            book.worksheets[0][f"{col}2"] = "=1+1"
    original = workbook(mutate)
    caches = {
        "R2": ("n", "2"), "S2": ("n", "0"), "U2": ("e", "#VALUE!"),
        "V2": ("b", "1"), "W2": ("n", "1.5"), "X2": ("n", "100"),
        "Y2": ("n", "0"), "AA2": ("n", "100.01"), "AB2": ("str", "2"),
        "AC2": ("n", "-1"),
    }
    output = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(output, "w") as target:
        for name in source.namelist():
            payload = source.read(name)
            if name == "xl/worksheets/sheet1.xml":
                root = ElementTree.fromstring(payload)
                ns = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
                for coordinate, (kind, value) in caches.items():
                    cell = root.find(f".//s:c[@r='{coordinate}']", ns)
                    cell.set("t", kind)
                    cell.find("s:v", ns).text = value
                payload = ElementTree.tostring(root)
            target.writestr(name, payload)
    row = read(output.getvalue())["sheets"][0]["rows"][0]
    expected = {
        "R2": (2, "positive", []), "S2": (0, "zero", []),
        "T2": (None, "invalid", ["formula_cache_missing"]),
        "U2": (None, "invalid", ["formula_cache_error"]),
        "V2": (None, "invalid", ["invalid_type"]),
        "W2": (None, "invalid", ["non_integer"]), "X2": (100, "positive", []),
        "Y2": (0, "zero", []), "Z2": (None, "invalid", ["formula_cache_missing"]),
        "AA2": (None, "invalid", ["above_100_percent"]),
        "AB2": (None, "invalid", ["invalid_type"]), "AC2": (None, "invalid", ["negative"]),
    }
    expected_caches = {"R2": 2, "S2": 0, "U2": "#VALUE!", "V2": True, "W2": 1.5,
                       "X2": 100, "Y2": 0, "AA2": 100.01, "AB2": "2", "AC2": -1}
    for sample in row["identificationEvidence"]:
        for key in ("uniquePeptides", "sequenceCoveragePercent"):
            field = sample[key]
            value, status, flags = expected[field["cell"]]
            assert field == {"cell": field["cell"], "value": value, "status": status, "flags": flags}
            source = next(c for c in row["cells"] if c["cell"] == field["cell"])
            assert source["value"] == "=1+1" and source["type"] == "f"
            assert source["cachedValue"] == expected_caches.get(field["cell"])
    literal_error = read(workbook(lambda b: b.worksheets[0].__setitem__("R2", "#VALUE!")))
    field = literal_error["sheets"][0]["rows"][0]["identificationEvidence"][0]["uniquePeptides"]
    assert field["status"] == "invalid" and field["flags"] == ["source_error"]


PINNED = Path(__file__).resolve().parents[1] / "data/interim/pxd036717" / enrichment.SOURCE_FILE


@pytest.mark.skipif(not PINNED.exists(), reason="requires owner-supplied pinned supplement")
def test_pinned_workbook_matches_independent_source_audit():
    doc = enrichment.read_workbook(PINNED.read_bytes())
    expected = [
        (383, [239, 198, 230], 196, [201]),
        (365, [200, 156, 186], 152, [53, 207, 225, 300]),
        (420, [309, 231, 287], 212, [149, 236, 349]),
        (457, [220, 156, 169], 144, [151, 241, 365]),
    ]
    for sheet, (rows, coverage, means, groups) in zip(doc["sheets"], expected, strict=True):
        audit = sheet["audit"]
        assert audit["rowCount"] == rows
        assert list(audit["validByReplicate"].values()) == coverage
        assert audit["completeMeans"] == means
        assert audit["multiAccessionRows"] == groups
        assert not audit["duplicateAccessions"]
        assert not any("disagrees" in flag for flag in audit["flags"])
    assert sum(s["audit"]["formulaCells"] for s in doc["sheets"]) == 11009
    assert {
        state: sum(s["audit"]["detectionCounts"][state] for s in doc["sheets"])
        for state in enrichment.DETECTION_STATES
    } == {"both": 2581, "bait_only": 784, "control_only": 188, "neither": 1322, "invalid": 0}
    cikb = doc["sheets"][1]["rows"][0]
    assert cikb["cells"][7]["value"] == cikb["cells"][13]["value"] == 0
    assert cikb["comparisons"][2]["reconstructedLog2Ratio"] is None
    assert "zero_published_bait_log2_without_positive_intensity" in cikb["comparisons"][2]["flags"]


@pytest.mark.skipif(not PINNED.exists(), reason="requires owner-supplied pinned supplement")
def test_pinned_identification_matches_independently_counted_source_cells():
    doc = enrichment.read_workbook(PINNED.read_bytes())
    # Counted directly from columns R:W and X:AC, independently of reader helpers.
    expected = [(1995, 303, 2000, 298, 1726, 743), (1729, 461, 1735, 455, 1217, 1055),
                (1966, 554, 1974, 546, 1886, 1237), (2001, 741, 2011, 731, 1305, 1461)]
    for table, (unique, zero_unique, coverage, zero_coverage, positive_lfq, below_minimum) in zip(
            doc["sheets"], expected, strict=True):
        assert table["audit"]["identificationEvidence"] == {
            "sampleCount": 6 * len(table["rows"]), "positiveLfqSamples": positive_lfq,
            "uniquePeptides": {"missing": 0, "zero": zero_unique, "positive": unique, "invalid": 0},
            "sequenceCoveragePercent": {"missing": 0, "zero": zero_coverage,
                                        "positive": coverage, "invalid": 0},
            "flags": {"below_reported_unique_peptide_minimum": below_minimum,
                      "spectral_counts_not_supplied": 6 * len(table["rows"])},
        }
    assert sum(s["audit"]["identificationEvidence"]["sampleCount"] for s in doc["sheets"]) == 9750


def test_cli_source_protection_pin_and_success(tmp_path, monkeypatch, capsys):
    source = tmp_path / "source.xlsx"
    source.write_bytes(workbook())
    with pytest.raises(SystemExit):
        enrichment.main(["--input", str(source), "--output", str(source)])
    hardlink = tmp_path / "hardlink.xlsx"
    hardlink.hardlink_to(source)
    source_before = source.read_bytes()
    with pytest.raises(SystemExit):
        enrichment.main(["--input", str(source), "--output", str(hardlink)])
    assert source.read_bytes() == source_before
    symlink = tmp_path / "symlink.xlsx"
    symlink.symlink_to(source)
    with pytest.raises(SystemExit):
        enrichment.main(["--input", str(source), "--output", str(symlink)])
    output = tmp_path / "new" / "evidence.json"
    args = ["--input", str(source), "--output", str(output)]
    with pytest.raises(ValueError, match="SHA-256"):
        enrichment.main(args)
    assert not output.exists()
    source_before = source.read_bytes()
    original_reader = enrichment.read_workbook
    monkeypatch.setattr(enrichment, "read_workbook", lambda data: original_reader(
        data, expected_sha256=hashlib.sha256(source_before).hexdigest()))
    assert enrichment.main(args) == 0
    assert source.read_bytes() == source_before
    assert json.loads(output.read_text())["sheets"][0]["audit"]["rowCount"] == 2
    assert json.loads(capsys.readouterr().out)["RpaA"]["validByReplicate"]["A"] == 1
