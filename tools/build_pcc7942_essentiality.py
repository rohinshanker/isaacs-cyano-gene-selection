#!/usr/bin/env python3
"""Build and verify the pinned PCC 7942 essentiality transfer dataset.

Usage:
    python3 tools/build_pcc7942_essentiality.py
    python3 tools/build_pcc7942_essentiality.py --check
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import tempfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook


WORKBOOK_PATH = Path("data/essentiality/source/mbio.00862-22-s0001.xlsx")
CROSSWALK_PATH = Path(
    "data/annotation/releases/GCF_000817325.1-RS_2026_05_13/"
    "identifier-crosswalk-v1.tsv"
)
GENES_PATH = Path("site/data/genes.json")
MANIFEST_PATH = Path("data/manifest/pcc7942-essentiality-v1.json")
OUTPUT_PATH = Path("site/data/pcc7942-essentiality-v1.json")

WORKBOOK_SHA256 = "b988b744c4c939ce6f47232eacfc30338907a9b911830999eb23414cbe6c331b"
WORKBOOK_SIZE = 1_359_396
WORKSHEET = "PG_metadata"
REQUIRED_COLUMNS = (
    "PG_ID",
    "UTEX 2973 NCBI",
    "PCC 7942 NCBI",
    "PCC 7942 essentiality",
)
SOURCE_STATUSES = (
    "essential",
    "beneficial",
    "non-essential",
    "ambiguous",
    "not_analyzed",
    "missing",
)
UNKNOWN_REASONS = (
    "source_unmatched",
    "source_multivalued",
    "source_pcc_missing",
    "crosswalk_unmatched",
    "crosswalk_multivalued",
    "crosswalk_ambiguous",
    "crosswalk_conflict",
)
MAPPING_STATUSES = ("accepted", "unmatched", "ambiguous", "conflicting")


class EssentialityDataError(ValueError):
    """Raised when a source or generated essentiality artifact is invalid."""


def sha256(path: Path) -> str:
    """Returns the lowercase SHA-256 digest of a file."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def verify_workbook(path: Path) -> None:
    """Checks the pinned workbook's exact byte identity."""
    if not path.is_file():
        raise EssentialityDataError(f"workbook is missing: {path}")
    if path.stat().st_size != WORKBOOK_SIZE:
        raise EssentialityDataError(
            f"workbook byte size mismatch: {path.stat().st_size} != {WORKBOOK_SIZE}"
        )
    actual = sha256(path)
    if actual != WORKBOOK_SHA256:
        raise EssentialityDataError(
            f"workbook SHA-256 mismatch: {actual} != {WORKBOOK_SHA256}"
        )


def _optional_text(value: Any, field: str) -> str | None:
    """Normalizes a nullable Excel text cell without accepting coercion."""
    if value is None:
        return None
    if not isinstance(value, str):
        raise EssentialityDataError(f"{field} must be text or empty, got {value!r}")
    if value != value.strip():
        raise EssentialityDataError(f"{field} has surrounding whitespace: {value!r}")
    return value or None


def read_workbook_rows(path: Path) -> list[dict[str, Any]]:
    """Reads the four source columns used by the exact join."""
    verify_workbook(path)
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if WORKSHEET not in workbook.sheetnames:
            raise EssentialityDataError(f"workbook lacks worksheet {WORKSHEET!r}")
        sheet = workbook[WORKSHEET]
        iterator = sheet.iter_rows(values_only=True)
        try:
            header = next(iterator)
        except StopIteration as error:
            raise EssentialityDataError("workbook worksheet is empty") from error
        named_header = [name for name in header if name is not None]
        if len(named_header) != len(set(named_header)):
            raise EssentialityDataError("workbook worksheet has duplicate columns")
        positions = {
            name: index for index, name in enumerate(header) if name is not None
        }
        missing = [name for name in REQUIRED_COLUMNS if name not in positions]
        if missing:
            raise EssentialityDataError(
                f"workbook worksheet lacks columns: {', '.join(missing)}"
            )

        rows: list[dict[str, Any]] = []
        pg_ids: set[int] = set()
        for excel_row, values in enumerate(iterator, start=2):
            raw_pg_id = values[positions["PG_ID"]]
            if isinstance(raw_pg_id, bool) or not isinstance(raw_pg_id, int):
                raise EssentialityDataError(
                    f"row {excel_row} PG_ID must be an integer, got {raw_pg_id!r}"
                )
            if raw_pg_id in pg_ids:
                raise EssentialityDataError(f"duplicate PG_ID: {raw_pg_id}")
            pg_ids.add(raw_pg_id)
            status = _optional_text(
                values[positions["PCC 7942 essentiality"]],
                "PCC 7942 essentiality",
            ) or "missing"
            if status not in SOURCE_STATUSES:
                raise EssentialityDataError(
                    f"row {excel_row} has unsupported essentiality status {status!r}"
                )
            rows.append(
                {
                    "pgId": raw_pg_id,
                    "utexLocus": _optional_text(
                        values[positions["UTEX 2973 NCBI"]], "UTEX 2973 NCBI"
                    ),
                    "pccLocus": _optional_text(
                        values[positions["PCC 7942 NCBI"]], "PCC 7942 NCBI"
                    ),
                    "status": status,
                }
            )
        return rows
    finally:
        workbook.close()


def read_gene_loci(path: Path) -> list[str]:
    """Returns sorted, unique locus tags from the plotted CDS dataset."""
    try:
        payload = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise EssentialityDataError(f"cannot load plotted genes from {path}: {error}") from error
    if not isinstance(payload, list):
        raise EssentialityDataError("plotted genes must be a JSON array")
    loci: list[str] = []
    for row in payload:
        if not isinstance(row, dict) or not isinstance(row.get("id"), str):
            raise EssentialityDataError("plotted gene row lacks a string id")
        loci.append(row["id"])
    if len(loci) != len(set(loci)):
        raise EssentialityDataError("plotted genes contain duplicate locus tags")
    return sorted(loci)


def read_crosswalk(path: Path) -> dict[str, list[dict[str, str]]]:
    """Loads only pinned PCC 7942 ortholog relationships from the crosswalk."""
    required = {
        "subject_locus_tag",
        "relationship",
        "object_namespace",
        "object_id",
        "mapping_ambiguity",
        "mapping_method",
    }
    grouped: dict[str, list[dict[str, str]]] = defaultdict(list)
    try:
        with path.open(encoding="utf-8", newline="") as handle:
            reader = csv.DictReader(handle, delimiter="\t")
            missing = required - set(reader.fieldnames or ())
            if missing:
                raise EssentialityDataError(
                    f"crosswalk lacks columns: {', '.join(sorted(missing))}"
                )
            for row in reader:
                if row["relationship"] != "pcc7942_ortholog":
                    continue
                if row["object_namespace"] != "PCC7942_RefSeq_locus_tag":
                    raise EssentialityDataError(
                        "PCC relationship has unexpected object namespace"
                    )
                if row["mapping_method"] != "exact shared RefSeq protein_id":
                    raise EssentialityDataError(
                        "PCC relationship is not an exact shared-protein mapping"
                    )
                grouped[row["subject_locus_tag"]].append(row)
    except OSError as error:
        raise EssentialityDataError(f"cannot read crosswalk {path}: {error}") from error
    return dict(grouped)


def join_by_locus(
    gene_loci: Iterable[str],
    source_rows: Iterable[dict[str, Any]],
    crosswalk: dict[str, list[dict[str, str]]],
) -> dict[str, dict[str, Any]]:
    """Transfers calls only when one exact source pair agrees with one crosswalk."""
    source_by_utex: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in source_rows:
        if row["utexLocus"] is not None:
            source_by_utex[row["utexLocus"]].append(row)

    by_locus: dict[str, dict[str, Any]] = {}
    for locus in sorted(gene_loci):
        candidates = source_by_utex.get(locus, [])
        links = crosswalk.get(locus, [])
        reason: str | None = None
        source: dict[str, Any] | None = None
        if not candidates:
            reason = "source_unmatched"
        elif len(candidates) != 1:
            reason = "source_multivalued"
        else:
            source = candidates[0]
            if source["pccLocus"] is None:
                reason = "source_pcc_missing"
            elif not links:
                reason = "crosswalk_unmatched"
            elif len(links) != 1:
                reason = "crosswalk_multivalued"
            elif links[0]["mapping_ambiguity"]:
                reason = "crosswalk_ambiguous"
            elif links[0]["object_id"] != source["pccLocus"]:
                reason = "crosswalk_conflict"

        if reason is None:
            assert source is not None
            by_locus[locus] = {
                "status": source["status"],
                "pccLocusTag": source["pccLocus"],
                "pangenomeId": source["pgId"],
                "mappingStatus": "accepted",
                "mappingReason": "exact_source_crosswalk_agreement",
            }
        else:
            if reason in {"source_unmatched", "source_pcc_missing", "crosswalk_unmatched"}:
                mapping_status = "unmatched"
            elif reason == "crosswalk_conflict":
                mapping_status = "conflicting"
            else:
                mapping_status = "ambiguous"
            by_locus[locus] = {
                "status": "unknown",
                "pccLocusTag": None,
                "pangenomeId": None,
                "mappingStatus": mapping_status,
                "mappingReason": reason,
            }
    return by_locus


def _load_manifest(path: Path) -> dict[str, Any]:
    try:
        manifest = json.loads(path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as error:
        raise EssentialityDataError(f"cannot load manifest {path}: {error}") from error
    if not isinstance(manifest, dict) or manifest.get("schemaVersion") != 1:
        raise EssentialityDataError("essentiality manifest schemaVersion must be 1")
    return manifest


def build_payload(root: Path) -> dict[str, Any]:
    """Builds the complete static payload from the three pinned inputs."""
    manifest = _load_manifest(root / MANIFEST_PATH)
    source_rows = read_workbook_rows(root / WORKBOOK_PATH)
    loci = read_gene_loci(root / GENES_PATH)
    by_locus = join_by_locus(
        loci,
        source_rows,
        read_crosswalk(root / CROSSWALK_PATH),
    )
    status_counts = Counter(row["status"] for row in by_locus.values())
    mapping_status_counts = Counter(
        row["mappingStatus"] for row in by_locus.values()
    )
    reason_counts = Counter(
        row["mappingReason"]
        for row in by_locus.values()
        if row["status"] == "unknown"
    )
    source_status_counts = Counter(row["status"] for row in source_rows)
    payload = {
        "schemaVersion": 1,
        "datasetId": "pcc7942-essentiality-transfer-v1",
        "source": manifest["source"],
        "joinPolicy": manifest["joinPolicy"],
        "counts": {
            "plottedUtex2973Loci": len(loci),
            "admitted": len(loci) - status_counts["unknown"],
            "unknown": status_counts["unknown"],
            "byStatus": {
                status: status_counts[status]
                for status in (*SOURCE_STATUSES, "unknown")
            },
            "byMappingStatus": {
                status: mapping_status_counts[status]
                for status in MAPPING_STATUSES
            },
            "unknownByReason": {
                reason: reason_counts[reason] for reason in UNKNOWN_REASONS
            },
            "workbookRows": len(source_rows),
            "workbookByStatus": {
                status: source_status_counts[status] for status in SOURCE_STATUSES
            },
        },
        "byLocus": by_locus,
    }
    validate_payload(payload, loci)
    return payload


def validate_payload(payload: dict[str, Any], gene_loci: Iterable[str]) -> None:
    """Validates ordering, schema, counts, and missing-value semantics."""
    by_locus = payload.get("byLocus")
    if not isinstance(by_locus, dict):
        raise EssentialityDataError("byLocus must be an object")
    expected_loci = sorted(gene_loci)
    actual_loci = list(by_locus)
    if actual_loci != expected_loci:
        raise EssentialityDataError(
            "byLocus must cover each plotted locus in sorted key order"
        )

    status_counts: Counter[str] = Counter()
    mapping_status_counts: Counter[str] = Counter()
    reason_counts: Counter[str] = Counter()
    expected_fields = {
        "status", "pccLocusTag", "pangenomeId", "mappingStatus", "mappingReason"
    }
    for row in by_locus.values():
        if set(row) != expected_fields:
            raise EssentialityDataError("record fields differ from schema")
        status = row["status"]
        status_counts[status] += 1
        mapping_status_counts[row["mappingStatus"]] += 1
        if status == "unknown":
            reason = row["mappingReason"]
            if reason not in UNKNOWN_REASONS:
                raise EssentialityDataError(f"invalid unknown reason: {reason!r}")
            reason_counts[reason] += 1
            expected_mapping_status = (
                "unmatched"
                if reason in {
                    "source_unmatched", "source_pcc_missing", "crosswalk_unmatched"
                }
                else "conflicting"
                if reason == "crosswalk_conflict"
                else "ambiguous"
            )
            if row["mappingStatus"] != expected_mapping_status:
                raise EssentialityDataError("unknown mappingStatus disagrees with reason")
            if row["pccLocusTag"] is not None or row["pangenomeId"] is not None:
                raise EssentialityDataError("unknown row must not expose an unadmitted join")
        else:
            if status not in SOURCE_STATUSES:
                raise EssentialityDataError(f"invalid admitted status: {status!r}")
            if row["mappingStatus"] != "accepted":
                raise EssentialityDataError("admitted row must have accepted mappingStatus")
            if row["mappingReason"] != "exact_source_crosswalk_agreement":
                raise EssentialityDataError("admitted row has an invalid mappingReason")
            if not isinstance(row["pccLocusTag"], str):
                raise EssentialityDataError("admitted row lacks a PCC 7942 locus")
            if (
                isinstance(row["pangenomeId"], bool)
                or not isinstance(row["pangenomeId"], int)
            ):
                raise EssentialityDataError("admitted row lacks an integer PG_ID")

    counts = payload.get("counts", {})
    expected_by_status = {
        status: status_counts[status] for status in (*SOURCE_STATUSES, "unknown")
    }
    expected_reasons = {reason: reason_counts[reason] for reason in UNKNOWN_REASONS}
    if counts.get("byStatus") != expected_by_status:
        raise EssentialityDataError("byStatus counts do not match records")
    expected_mapping_statuses = {
        status: mapping_status_counts[status] for status in MAPPING_STATUSES
    }
    if counts.get("byMappingStatus") != expected_mapping_statuses:
        raise EssentialityDataError("byMappingStatus counts do not match records")
    if counts.get("unknownByReason") != expected_reasons:
        raise EssentialityDataError("unknownByReason counts do not match records")
    if counts.get("admitted") != len(by_locus) - status_counts["unknown"]:
        raise EssentialityDataError("admitted count does not match records")
    if counts.get("unknown") != status_counts["unknown"]:
        raise EssentialityDataError("unknown count does not match records")


def render_payload(payload: dict[str, Any]) -> str:
    """Renders stable UTF-8 JSON with a final newline."""
    return json.dumps(payload, ensure_ascii=False, indent=2, sort_keys=True) + "\n"


def write_or_check(root: Path, check: bool) -> None:
    """Writes the artifact or proves it matches a fresh temporary rebuild."""
    rendered = render_payload(build_payload(root))
    output = root / OUTPUT_PATH
    if check:
        with tempfile.TemporaryDirectory() as directory:
            candidate = Path(directory) / output.name
            candidate.write_text(rendered, encoding="utf-8")
            try:
                current = output.read_bytes()
            except OSError as error:
                raise EssentialityDataError(f"cannot read generated output: {error}") from error
            if candidate.read_bytes() != current:
                raise EssentialityDataError(
                    f"{OUTPUT_PATH} differs from a deterministic rebuild"
                )
        return
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(rendered, encoding="utf-8")


def main() -> int:
    """Runs the command-line builder."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--check",
        action="store_true",
        help="compare the checked-in artifact with a fresh deterministic rebuild",
    )
    parser.add_argument(
        "--root",
        type=Path,
        default=Path(__file__).resolve().parents[1],
        help=argparse.SUPPRESS,
    )
    args = parser.parse_args()
    try:
        write_or_check(args.root.resolve(), args.check)
    except EssentialityDataError as error:
        parser.error(str(error))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
