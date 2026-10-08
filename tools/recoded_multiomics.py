#!/usr/bin/env python3
"""Read Nyerges Supplementary Data 2/3 without merging strains or scales.

The output preserves source sheet, row and column identities. It is an intake
artifact, not a gene join: admitting it requires an explicit strain/genome and
identifier crosswalk. Missing observations never become zero. Published
no-growth sentinels are retained as evidence but never as doubling times.
"""

from __future__ import annotations

import argparse
import json
import math
from itertools import chain
from pathlib import Path
from typing import Any, Iterable, Sequence

import openpyxl

from fetch_recoded_ecoli import DEFAULT_DEST, PINNED, verify


FITNESS_HEADERS = (
    "Strain", "Doubling time (minutes)",
    "Doubling time Standard Deviation (minutes)", "Maximum OD600",
    "Maximum OD600 Standard Deviation",
    *(f"DT Replicate {i}" for i in range(1, 11)),
    *(f"OD600 Replicate {i}" for i in range(1, 11)),
)
BIOLOG_HEADERS = ("Plate", "Well", "Substrate", "Max Height")
OMICS_SHEETS = (
    "Seg9-18_36-44_46-49_51-59", "Seg30-35", "Seg30-35_debugged",
    "Seg70-81", "Syn61_delta3_ev5",
)
DERIVED_COLUMNS = {
    "RNA_LFC": ("transcript_log2_fold_change", "log2 fold change, read-count calculation"),
    "RNA_LFC_EdgeR": ("transcript_log2_fold_change_edger", "log2 fold change, EdgeR"),
    "RNA_P-value": ("transcript_p_value", "p-value"),
    "RIBO_LFC": ("ribosome_log2_fold_change", "log2 fold change, read-count calculation"),
    "RIBO_LFC_EdgeR": ("ribosome_log2_fold_change_edger", "log2 fold change, EdgeR"),
    "RIBO_P-value": ("ribosome_p_value", "p-value"),
    "Delta_LFC": ("translation_efficiency_log2_fold_change", "log2 fold change"),
}
NO_GROWTH = "0 (no growth detected)"


def blank(value: Any) -> bool:
    """Whether a cell is blank, including whitespace-only worksheet padding."""
    return value is None or isinstance(value, str) and not value.strip()


def number(value: Any, where: str, *, nonnegative: bool = False) -> float | None:
    """Read a finite source value; a hyphen or blank is absent, not zero."""
    if blank(value) or value == "-":
        return None
    if isinstance(value, bool):
        raise ValueError(f"{where}: Boolean is not a measurement")
    try:
        result = float(value)
    except (TypeError, ValueError) as error:
        raise ValueError(f"{where}: invalid measurement {value!r}") from error
    if not math.isfinite(result) or nonnegative and result < 0:
        raise ValueError(f"{where}: invalid measurement {value!r}")
    return result


def text_cell(value: Any, where: str) -> str:
    """Require a nonempty source identifier without changing its namespace."""
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f"{where}: missing text identifier")
    return value.strip()


def checked_rows(rows: Iterable[Sequence[Any]], headers: Sequence[str], sheet: str):
    """Validate a table's exact header and yield numbered nonempty rows."""
    iterator = iter(rows)
    if tuple(next(iterator, ())) != tuple(headers):
        raise ValueError(f"{sheet}: unexpected columns")
    for row_number, row in enumerate(iterator, 2):
        if all(blank(cell) for cell in row):
            continue
        if len(row) != len(headers):
            raise ValueError(f"{sheet}:{row_number}: wrong column count")
        yield row_number, row


def read_growth(rows: Iterable[Sequence[Any]]) -> list[dict[str, Any]]:
    """Read strain-condition growth records and their separate replicates."""
    result = []
    seen = set()
    for row_number, row in checked_rows(rows, FITNESS_HEADERS, "Fitness_Source_data"):
        label = text_cell(row[0], f"Fitness_Source_data:A{row_number}")
        if label in seen:
            raise ValueError(f"Duplicate growth source label: {label}")
        seen.add(label)
        no_growth = row[1] == NO_GROWTH
        values = [number(v, f"Fitness_Source_data:{row_number}:{i + 1}", nonnegative=True)
                  for i, v in enumerate(row) if i > 0 and not (i == 1 and no_growth)]
        if no_growth:
            # The source uses zero as a no-growth sentinel for these cells.
            # Keep the original cells separately, never expose 0 min as a rate.
            values.insert(0, None)
            if any(v not in (None, 0) for v in values):
                raise ValueError(f"{label}: no-growth sentinel conflicts with measurements")
        elif values[0] is not None and values[0] <= 0:
            raise ValueError(f"{label}: doubling time must be positive")
        result.append({
            "sourceRow": row_number, "sourceLabel": label,
            "growthStatus": "no_growth_detected" if no_growth else "reported",
            "doublingTimeMinutes": None if no_growth else values[0],
            "doublingTimeSdMinutes": None if no_growth else values[1],
            "maximumOd600": None if no_growth else values[2],
            "maximumOd600Sd": None if no_growth else values[3],
            "doublingTimeReplicatesMinutes": [None] * 10 if no_growth else values[4:14],
            "maximumOd600Replicates": [None] * 10 if no_growth else values[14:24],
            "sourceCells": list(row),
        })
    return result


def read_biolog(rows: Iterable[Sequence[Any]], sheet: str) -> list[dict[str, Any]]:
    """Unpivot plate blocks; do not reinterpret the deposited Max Height scale."""
    headers = tuple(value for i in range(5)
                    for value in ((*BIOLOG_HEADERS, None) if i < 4 else BIOLOG_HEADERS))
    result = []
    seen = set()
    for row_number, row in checked_rows(rows, headers, sheet):
        for offset in (0, 5, 10, 15, 20):
            block = row[offset:offset + 4]
            if all(blank(value) for value in block):
                continue
            plate, well, substrate = (text_cell(block[i], f"{sheet}:{row_number}:{offset+i+1}")
                                      for i in range(3))
            if plate not in {"PM01", "PM02", "PM04", "PM06", "PM09"}:
                raise ValueError(f"{sheet}: unexpected plate {plate}")
            if (plate, well) in seen:
                raise ValueError(f"{sheet}: duplicate well {plate}/{well}")
            seen.add((plate, well))
            result.append({"sourceRow": row_number, "sourceColumn": offset + 4,
                           "plate": plate, "well": well, "substrate": substrate,
                           "maxHeight": number(block[3], f"{sheet}:{row_number}:{offset+4}")})
    return result


def omics_column(header: str) -> dict[str, str]:
    """Decode the publisher's explicit column grammar, keeping its full label."""
    if header in DERIVED_COLUMNS:
        kind, units = DERIVED_COLUMNS[header]
        return {"sourceColumn": header, "kind": kind, "units": units}
    for prefix, assay in (("RNA-", "transcript"), ("RIBO-", "ribosome"),
                          ("RIBOE_coli_DH10B_", "ribosome")):
        # The final prefix is a literal, pinned publisher header typo.
        if header.startswith(prefix):
            if header.endswith("_rpkm"):
                return {"sourceColumn": header, "kind": f"{assay}_rpkm", "units": "RPKM"}
            if header.endswith("_reads"):
                return {"sourceColumn": header, "kind": f"{assay}_read_count", "units": "reads"}
    raise ValueError(f"Unrecognized omics column: {header}")


def read_omics(rows: Iterable[Sequence[Any]], sheet: str) -> dict[str, Any]:
    """Read one contrast without combining control columns across contrasts."""
    iterator = iter(rows)
    headers = tuple(next(iterator, ()))
    if headers[:2] != ("Gene", "Description") or len(set(headers)) != len(headers):
        raise ValueError(f"{sheet}: invalid gene/measurement columns")
    if not set(DERIVED_COLUMNS) <= set(headers):
        raise ValueError(f"{sheet}: incomplete derived columns")
    columns = [omics_column(header) for header in headers[2:]]
    records = []
    seen = set()
    for row_number, row in checked_rows(chain((headers,), iterator), headers, sheet):
        gene = text_cell(row[0], f"{sheet}:A{row_number}")
        if gene in seen:
            raise ValueError(f"{sheet}: duplicate gene identifier {gene}")
        seen.add(gene)
        values = []
        missing = {}
        for index, (column, cell) in enumerate(zip(columns, row[2:], strict=True)):
            value = number(cell, f"{sheet}:{row_number}:{index+3}",
                           nonnegative=column["units"] in {"RPKM", "reads", "p-value"})
            if value is not None and column["units"] == "p-value" and value > 1:
                raise ValueError(f"{sheet}:{row_number}: p-value exceeds one")
            if value is not None and column["units"] == "reads" and not value.is_integer():
                raise ValueError(f"{sheet}:{row_number}: fractional read count")
            if value is None:
                missing[str(index)] = "hyphen" if cell == "-" else "blank"
            values.append(value)
        records.append({"sourceRow": row_number, "gene": gene,
                        "description": row[1], "values": values, "missing": missing})
    return {"sheet": sheet, "identifierNamespace": "published Gene column",
            "columns": columns, "rows": records}


def extract(source_dir: Path) -> dict[str, Any]:
    """Verify the pinned workbooks and extract literal, auditable source tables."""
    sources = []
    for name in ("Supplementary_Data_2.xlsx", "Supplementary_Data_3.xlsx"):
        pinned = next(item for item in PINNED if item.name == name)
        problems = verify(source_dir / name, pinned)
        if problems:
            raise ValueError("; ".join(problems))
        sources.append({"file": name, "sha256": pinned.sha256, "bytes": pinned.size})
    with (source_dir / sources[0]["file"]).open("rb") as handle:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
        try:
            growth = read_growth(workbook["Fitness_Source_data"].values)
            biolog = [{"sheet": sheet, "rows": read_biolog(workbook[sheet].values, sheet)}
                      for sheet in workbook.sheetnames
                      if sheet not in {"Fitness_Source_data", "Biolog_Source_data"}]
        finally:
            workbook.close()
    with (source_dir / sources[1]["file"]).open("rb") as handle:
        workbook = openpyxl.load_workbook(handle, read_only=True, data_only=True)
        try:
            if set(workbook.sheetnames) != {"Legend", *OMICS_SHEETS}:
                raise ValueError("Supplementary Data 3: unexpected sheets")
            omics = [read_omics(workbook[sheet].values, sheet) for sheet in OMICS_SHEETS]
        finally:
            workbook.close()
    return {"schemaVersion": 1, "studyId": "nyerges-2026", "admission": "unjoined-source-tables",
            "sources": sources, "growth": growth, "biolog": biolog, "omics": omics}


def main(argv: list[str] | None = None) -> int:
    """Write deterministic intake JSON outside published organism directories."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=DEFAULT_DEST)
    parser.add_argument("--output", type=Path,
                        default=Path("data/interim/recoded-ecoli/source-tables.json"))
    args = parser.parse_args(argv)
    document = extract(args.source_dir)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(document, ensure_ascii=False, allow_nan=False,
                                      separators=(",", ":")) + "\n", encoding="utf-8")
    print(f"Wrote {len(document['growth'])} growth records, "
          f"{sum(len(table['rows']) for table in document['biolog'])} Biolog wells and "
          f"{len(document['omics'])} separate omics contrasts to {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
