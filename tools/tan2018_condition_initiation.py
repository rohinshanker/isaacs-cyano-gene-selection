#!/usr/bin/env python3
"""Build the four Tan 2018 condition TSS-initiation layers from Table S1.

Usage:
    tan2018_condition_initiation.py fetch [--workbook PATH]
    tan2018_condition_initiation.py write [--genes PATH] [--out-dir PATH]
    tan2018_condition_initiation.py check [--genes PATH] [--out-dir PATH]
    tan2018_condition_initiation.py verify-source [--workbook PATH] [--article PATH]
    tan2018_condition_initiation.py summary [--genes PATH]

Owner decision, 2026-10-12: "use the s1 table". Each layer is, per current
gene, the sum over the gene's Table S1 gene-associated TSS (``gTSS``) rows of
the mean of that condition's two biological replicates:

* control: ``CT1_RawReads``, ``CT2_RawReads``
* dark: ``D1_RawReads``, ``D2_RawReads``
* high light: ``HL1_RawReads``, ``HL2_RawReads``
* high temperature: ``HT1_RawReads``, ``HT2_RawReads``

A row counts for a gene only when its published ``Locus_tags`` value is exactly
that gene's current ``M744_RS`` locus tag; nothing is matched by symbol,
coordinate or distance. Published counts are used as printed, fractional ones
included, so a value may end in .25 or .5. A gene with mapped rows whose counts
are all zero scores 0; a gene with no mapped row has no table row and stays
unknown. The legacy pooled ``TAN2018_TSS`` table uses a different, recovered
site set (``tools/tan2018_pooled_audit.py``) and is neither read nor changed;
these layers do not average to it.

The input is the pinned gTSS extract of Table S1
(``scripts/prepare_tan2018.py``). ``verify-source`` re-extracts it from the
checksum-pinned publisher workbook and re-matches the article's licence lines.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import io
import json
import re
import sys
from collections import Counter
from dataclasses import dataclass
from decimal import Decimal
from fractions import Fraction
from pathlib import Path
from typing import Any, Collection, Mapping, Sequence

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from scripts import prepare_tan2018  # noqa: E402
from tools import ingest_expression as ingest  # noqa: E402
from tools import tan2018_pooled_audit as pooled_audit  # noqa: E402

EXPRESSION = ROOT / "data/expression"
GENES = ROOT / "site/data/genes.json"
SOURCES = EXPRESSION / "sources.json"
EXTRACT = EXPRESSION / "tan2018_utex2973_tss_table_s1.tsv"
EXTRACT_SHA256 = "d06524cf492aa8f009bf004c596a5d305ac54acbe7537ec98bbc7992ce164bff"
EXTRACT_ROWS = 2475
WORKBOOK_URL = ("https://static-content.springer.com/esm/art%3A10.1186%2Fs13068-018-1215-8/"
                "MediaObjects/13068_2018_1215_MOESM1_ESM.xlsx")
WORKBOOK = pooled_audit.INTERIM / "13068_2018_1215_MOESM1_ESM.xlsx"
WORKBOOK_SHA256 = prepare_tan2018.SOURCE_SHA256
# Table S1's header row as published, in order, without its trailing empty cells.
WORKBOOK_HEADER = (
    "TSS_ID", "Chromosome", "Type", "Strand", "TSS_position", "TssCovRatio",
    "LocalCoverageEnrichmentScore", "LocalTssEnrichmentScore", "TSS_nt",
    "Distances to start codon", "Locus_tags", "Gene Symbols", "Mean_Norm_Reads",
    "Max_Norm_Reads", "DK_vs_CT_log2FoldChange", "DK_vs_CT_padj",
    "HL_vs_CT_log2FoldChange", "HL_vs_CT_padj", "HT_vs_CT_log2FoldChange",
    "HT_vs_CT_padj", "CT1_RawReads", "CT2_RawReads", "D1_RawReads", "D2_RawReads",
    "HL1_RawReads", "HL2_RawReads", "HT1_RawReads", "HT2_RawReads", "Product",
    "Description",
)
# The article's licence lines: the article is CC BY 4.0, and the publisher adds
# a CC0 waiver for "the data made available in this article, unless otherwise stated".
LICENCE_QUOTES = (
    "Open Access This article is distributed under the terms of the Creative Commons "
    "Attribution 4.0 International License",
    "The Creative Commons Public Domain Dedication waiver "
    "( http://creativecommons.org/publicdomain/zero/1.0/ ) applies to the data made "
    "available in this article, unless otherwise stated.",
)

CURRENT_LOCUS = re.compile(r"M744_RS\d{5}")
# A published count: a non-negative decimal of bounded length, never blank.
COUNT = re.compile(r"\d{1,9}(?:\.\d{1,6})?")
TABLE_HEADER = "locus_tag\tabundance\tsource_gene_id\n"


@dataclass(frozen=True)
class Condition:
    """One Tan 2018 condition layer and the Table S1 replicate columns it reads."""

    name: str
    source_id: str
    metric_key: str
    file: str
    replicates: tuple[str, str]
    columns: tuple[str, str]

    @property
    def workbook_columns(self) -> tuple[str, str]:
        """Table S1's own headers for the two replicates."""
        return tuple(f"{replicate}_RawReads" for replicate in self.replicates)  # type: ignore[return-value]


CONDITIONS = (
    Condition("control", "TAN2018_TSS_CONTROL", "tssInitiationControl",
              "tan2018_utex2973_tss_initiation_control.tsv",
              ("CT1", "CT2"), ("control_1", "control_2")),
    Condition("dark", "TAN2018_TSS_DARK", "tssInitiationDark",
              "tan2018_utex2973_tss_initiation_dark.tsv",
              ("D1", "D2"), ("dark_1", "dark_2")),
    Condition("high_light", "TAN2018_TSS_HIGH_LIGHT", "tssInitiationHighLight",
              "tan2018_utex2973_tss_initiation_high_light.tsv",
              ("HL1", "HL2"), ("high_light_1", "high_light_2")),
    Condition("high_temperature", "TAN2018_TSS_HIGH_TEMPERATURE",
              "tssInitiationHighTemperature",
              "tan2018_utex2973_tss_initiation_high_temperature.tsv",
              ("HT1", "HT2"), ("high_temperature_1", "high_temperature_2")),
)


@dataclass(frozen=True)
class Row:
    """One Table S1 gTSS row: its id, position, source locus and replicate counts."""

    tss_id: str
    position: int
    locus: str
    counts: Mapping[str, Fraction]


def require(condition: bool, message: str) -> None:
    """Raises a descriptive input-contract error."""
    if not condition:
        raise ValueError(message)


def sha256(data: bytes) -> str:
    """The lower-case SHA-256 of ``data``."""
    return hashlib.sha256(data).hexdigest()


def read_extract(data: bytes, expected_sha256: str = EXTRACT_SHA256,
                 expected_rows: int = EXTRACT_ROWS) -> list[Row]:
    """The pinned gTSS extract, refused unless every line is one well-formed row."""
    observed = sha256(data)
    require(observed == expected_sha256,
            f"Table S1 extract checksum mismatch: expected {expected_sha256}, got {observed}")
    text = data.decode("utf-8")
    require("\r" not in text and text.endswith("\n"),
            "Table S1 extract must be LF-terminated lines")
    lines = text[:-1].split("\n")
    header = lines[0].split("\t")
    require(header == list(prepare_tan2018.OUTPUT_COLUMNS),
            f"Table S1 extract header is {header}, not {list(prepare_tan2018.OUTPUT_COLUMNS)}")
    count_columns = [column for condition in CONDITIONS for column in condition.columns]
    rows: list[Row] = []
    seen: set[str] = set()
    for number, line in enumerate(lines[1:], start=2):
        cells = line.split("\t")
        require(len(cells) == len(header),
                f"Table S1 extract line {number} has {len(cells)} fields, not {len(header)}")
        row = dict(zip(header, cells))
        require(row["type"] == "gTSS",
                f"Table S1 extract line {number} is {row['type']!r}, not gTSS")
        require(row["tss_id"] not in seen,
                f"Table S1 extract line {number} repeats {row['tss_id']}")
        seen.add(row["tss_id"])
        require(row["position"].isdigit(),
                f"Table S1 extract line {number} has position {row['position']!r}")
        for column in count_columns:
            require(COUNT.fullmatch(row[column]) is not None,
                    f"Table S1 extract line {number} {column} must be a non-negative "
                    f"count, not {row[column]!r}")
        rows.append(Row(row["tss_id"], int(row["position"]), row["locus_tag"],
                        {c: Fraction(Decimal(row[c])) for c in count_columns}))
    require(len(rows) == expected_rows,
            f"Table S1 extract has {len(rows)} gTSS rows, not {expected_rows}")
    return rows


def mapping_reason(locus: str, targets: Collection[str]) -> str:
    """Why a row does or does not join a current gene, by exact locus tag only."""
    if locus == "":
        return "source_locus_missing"
    if CURRENT_LOCUS.fullmatch(locus) is None:
        return "source_locus_not_current_id"
    if locus not in targets:
        return "source_locus_absent_from_current_cds"
    return "exact_current_locus_id"


def associate(rows: Sequence[Row], targets: Collection[str]
              ) -> tuple[dict[str, list[Row]], Counter[str]]:
    """Each current gene's mapped rows, and the count of rows by mapping reason."""
    by_locus: dict[str, list[Row]] = {}
    reasons: Counter[str] = Counter()
    for row in rows:
        reason = mapping_reason(row.locus, targets)
        reasons[reason] += 1
        if reason == "exact_current_locus_id":
            by_locus.setdefault(row.locus, []).append(row)
    return dict(sorted(by_locus.items())), reasons


def condition_value(rows: Sequence[Row], condition: Condition) -> Fraction:
    """The sum over ``rows`` of the mean of the condition's two replicates."""
    first, second = condition.columns
    return sum(((row.counts[first] + row.counts[second]) / 2 for row in rows), Fraction(0))


def decimal_text(value: Fraction) -> str:
    """``value`` at six decimals, refusing any value that does not fit exactly."""
    return pooled_audit.decimal_text(value)


def layer_table(by_locus: Mapping[str, Sequence[Row]], condition: Condition) -> bytes:
    """One condition's table: locus, value, and the Table S1 TSS ids summed."""
    lines = [TABLE_HEADER]
    for locus, rows in sorted(by_locus.items()):
        ids = ";".join(row.tss_id for row in sorted(rows, key=lambda r: (r.position, r.tss_id)))
        lines.append(f"{locus}\t{decimal_text(condition_value(rows, condition))}\t{ids}\n")
    return "".join(lines).encode("ascii")


def build(extract: bytes, targets: Collection[str]) -> dict[str, bytes]:
    """Every condition's table, keyed by file name."""
    by_locus, _ = associate(read_extract(extract), targets)
    return {condition.file: layer_table(by_locus, condition) for condition in CONDITIONS}


def check_tables(tables: Mapping[str, bytes], out_dir: Path) -> None:
    """Refuse unless each committed table equals the rebuild."""
    for name, data in tables.items():
        path = out_dir / name
        require(path.is_file(), f"{path} is missing; run `write` first")
        require(path.read_bytes() == data, f"{path} differs from the rebuild")


def check_manifest(sources: Sequence[Mapping[str, Any]], tables: Mapping[str, bytes]) -> None:
    """Refuse unless sources.json declares each layer with its rebuilt checksum."""
    by_id = {source.get("id"): source for source in sources}
    for condition in CONDITIONS:
        source = by_id.get(condition.source_id)
        if source is None:
            raise ValueError(f"sources.json lacks {condition.source_id}")
        expected = {
            "file": condition.file,
            "metricKey": condition.metric_key,
            "sha256": sha256(tables[condition.file]),
            "quantity": "tss_initiation",
        }
        for key, value in expected.items():
            require(source.get(key) == value,
                    f"{condition.source_id} {key} is {source.get(key)!r}, not {value!r}")


def verify_workbook(workbook: Path, extract: bytes,
                    expected_sha256: str = WORKBOOK_SHA256) -> None:
    """Refuse unless the pinned workbook's Table S1 header and gTSS rows are the extract's."""
    from openpyxl import load_workbook  # Imported here: only this mode reads XLSX.

    require(workbook.is_file(),
            f"{workbook} is missing; run `tan2018_condition_initiation.py fetch` first")
    observed = sha256(workbook.read_bytes())
    require(observed == expected_sha256,
            f"Table S1 workbook checksum mismatch: expected {expected_sha256}, got {observed}")
    sheet = load_workbook(workbook, read_only=True, data_only=True)["Table S1"]
    header = list(list(sheet.iter_rows(min_row=2, max_row=2, values_only=True))[0])
    while header and header[-1] is None:
        header.pop()
    require(tuple(header) == WORKBOOK_HEADER,
            f"Table S1 header is {header}, not {WORKBOOK_HEADER}")
    rebuilt = io.StringIO()
    writer = csv.writer(rebuilt, delimiter="\t", lineterminator="\n")  # As write_table.
    writer.writerow(prepare_tan2018.OUTPUT_COLUMNS)
    writer.writerows(prepare_tan2018.extract_rows(workbook))
    require(rebuilt.getvalue().encode("utf-8") == extract,
            "the workbook's gTSS rows do not reproduce the pinned extract")


def verify_article(article: Path, expected_sha256: str = pooled_audit.ARTICLE_SHA256) -> None:
    """Refuse unless the pinned article carries both licence lines verbatim."""
    require(article.is_file(),
            f"{article} is missing; run `tan2018_pooled_audit.py fetch` first")
    text = pooled_audit.article_text(article.read_bytes(), expected_sha256)
    for quote in LICENCE_QUOTES:
        require(pooled_audit.collapse(quote) in text, f"the article lacks: {quote}")


def summary(extract: bytes, targets: Sequence[str]) -> dict[str, Any]:
    """Coverage, mapping and numeric examples, recomputed from the extract."""
    rows = read_extract(extract)
    by_locus, reasons = associate(rows, targets)
    layers = {}
    for condition in CONDITIONS:
        values = {locus: condition_value(group, condition) for locus, group in by_locus.items()}
        layers[condition.source_id] = {
            "genes": len(values),
            "zeroGenes": sum(1 for v in values.values() if v == 0),
            "fractionalGenes": sum(1 for v in values.values() if v.denominator != 1),
            "zeroRows": sum(1 for group in by_locus.values() for row in group
                            if condition_value([row], condition) == 0),
            "min": decimal_text(min(values.values())),
            "max": decimal_text(max(values.values())),
        }
    return {
        "rows": len(rows),
        "rowsByMappingReason": dict(sorted(reasons.items())),
        "genesWithValue": len(by_locus),
        "targetsWithoutValue": len(set(targets) - set(by_locus)),
        "maxRowsPerGene": max(len(group) for group in by_locus.values()),
        "layers": layers,
    }


def main(argv: Sequence[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Tan 2018 Table S1 condition layers.")
    parser.add_argument("mode", choices=("fetch", "write", "check", "verify-source", "summary"))
    parser.add_argument("--workbook", type=Path, default=WORKBOOK)
    parser.add_argument("--article", type=Path, default=pooled_audit.ARTICLE)
    parser.add_argument("--genes", type=Path, default=GENES)
    parser.add_argument("--extract", type=Path, default=EXTRACT)
    parser.add_argument("--out-dir", type=Path, default=EXPRESSION)
    parser.add_argument("--sources", type=Path, default=SOURCES)
    arguments = parser.parse_args(argv)
    try:
        if arguments.mode == "fetch":
            ingest.fetch(WORKBOOK_URL, WORKBOOK_SHA256, arguments.workbook)
            return 0
        extract = arguments.extract.read_bytes()
        if arguments.mode == "verify-source":
            verify_workbook(arguments.workbook, extract)
            verify_article(arguments.article)
            return 0
        targets = pooled_audit.target_loci(arguments.genes)
        if arguments.mode == "summary":
            print(json.dumps(summary(extract, targets), indent=2))
            return 0
        tables = build(extract, targets)
        if arguments.mode == "write":
            for name, data in tables.items():
                (arguments.out_dir / name).write_bytes(data)
            return 0
        check_tables(tables, arguments.out_dir)
        check_manifest(json.loads(arguments.sources.read_text(encoding="utf-8")), tables)
        return 0
    except ValueError as error:
        print(f"tan2018_condition_initiation: {error}", file=sys.stderr)
        return 1


if __name__ == "__main__":
    sys.exit(main())
