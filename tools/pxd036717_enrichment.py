"""Preserve the PXD036717 supplement and reconstruct non-imputed LFQ ratios.

This is a non-release reader, not a gene mapper or interactor classifier. LFQ
intensities are MaxQuant outputs, not raw mass spectra. Published formulas and
their cached values remain separate from independently reconstructed values.
"""

from __future__ import annotations

import argparse
from collections import Counter
from contextlib import ExitStack
import hashlib
import io
import json
import math
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


SOURCE_FILE = "1-s2.0-S1096717623001490-mmc1.xlsx"
SOURCE_SHA256 = "287d5f49a71425ba2b085a7070053a34f9f17219833663e1231ea3ab81d86804"
BAITS = ("RpaA", "CikB", "NblS", "ManS")
REPLICATES = ("A", "B", "C")
DETECTION_STATES = ("both", "bait_only", "control_only", "neither", "invalid")
ABS_TOLERANCE = 1e-7


def expected_headers(bait: str) -> list[str]:
    """Exact deposited headers, including their typo and trailing spaces."""
    headers = ["Uniprot indentifier", "Gene name"]
    for prefix in ("LFQ intensity", "log2 LFQ intensity"):
        for arm in ("control", f"{bait}-APEX2"):
            headers.extend(f"{prefix} {arm} {rep}" for rep in REPLICATES)
    headers.extend(
        f"log2 of LFQ ratios ({bait}-APEX2 {rep}/control {rep}) "
        for rep in REPLICATES
    )
    for prefix, suffix in (("Unique peptides", ""), ("Sequence coverage", " [%]")):
        for arm in ("control", f"{bait}-APEX2"):
            headers.extend(f"{prefix} {arm} {rep}{suffix}" for rep in REPLICATES)
    return headers + ["Mol. weight [kDa]", "Sequence length", "Fasta headers"]


def _finite_number(value: Any) -> bool:
    return type(value) in (int, float) and math.isfinite(value)


def reconstruct_comparison(
    control: Any, bait: Any, published_control_log2: Any,
    published_bait_log2: Any, published_ratio: Any,
) -> dict[str, Any]:
    """Compute log2(bait)-log2(control) only for two finite positive inputs.

    Source zero is unavailable intensity, never an observed log2 zero. The
    source's A/A, B/B, C/C indexing is preserved; this asserts no paired design.
    Diagnostics compare source cells without repairing or trusting their caches.
    """
    flags = []
    logs = []
    invalid_intensity = False
    for arm, raw, published_log in (
        ("control", control, published_control_log2),
        ("bait", bait, published_bait_log2),
    ):
        if raw is None or (_finite_number(raw) and raw == 0):
            flags.append(f"missing_{arm}_intensity")
        elif not _finite_number(raw) or raw < 0:
            flags.append(f"invalid_{arm}_intensity")
            invalid_intensity = True
        raw_log = math.log2(raw) if _finite_number(raw) and raw > 0 else None
        logs.append(raw_log)
        if not _finite_number(published_log):
            flags.append(f"unavailable_published_{arm}_log2")
        elif raw_log is None:
            flags.append(f"published_{arm}_log2_without_positive_intensity")
            if published_log == 0:
                flags.append(f"zero_published_{arm}_log2_without_positive_intensity")
        elif not math.isclose(published_log, raw_log, rel_tol=0, abs_tol=ABS_TOLERANCE):
            flags.append(f"published_{arm}_log2_disagrees_with_intensity")
    reconstructed = logs[1] - logs[0] if all(v is not None for v in logs) else None
    if not _finite_number(published_ratio):
        flags.append("unavailable_published_ratio")
    else:
        if all(_finite_number(v) for v in (published_control_log2, published_bait_log2)):
            expected = published_bait_log2 - published_control_log2
            if not math.isclose(published_ratio, expected, rel_tol=0, abs_tol=ABS_TOLERANCE):
                flags.append("published_ratio_disagrees_with_published_logs")
        if reconstructed is not None and not math.isclose(
            published_ratio, reconstructed, rel_tol=0, abs_tol=ABS_TOLERANCE
        ):
            flags.append("published_ratio_disagrees_with_reconstruction")
    if invalid_intensity:
        detection = "invalid"
    elif logs[0] is not None:
        detection = "both" if logs[1] is not None else "control_only"
    else:
        detection = "bait_only" if logs[1] is not None else "neither"
    return {
        "reconstructedLog2Ratio": reconstructed,
        "detectionState": detection, "flags": flags,
    }


def _source_cell(cell: Any, cached: Any, coordinate: str) -> dict[str, Any]:
    """Keep typed source and cached values; never evaluate Excel formulas."""
    for value in (cell.value, cached.value):
        if value is not None and type(value) not in (str, int, float, bool):
            raise ValueError(f"{coordinate}: unsupported source value type")
        if type(value) in (int, float) and not math.isfinite(value):
            raise ValueError(f"{coordinate}: non-finite source value")
    return {
        "cell": coordinate, "value": cell.value, "type": cell.data_type,
        "cachedValue": cached.value, "cachedType": cached.data_type,
    }


def read_workbook(data: bytes, *, expected_sha256: str = SOURCE_SHA256) -> dict[str, Any]:
    """Read a pinned workbook, retaining every populated data cell and formula.

    A different explicit pin is supported for synthetic tests or separately
    reviewed source revisions; the command line always requires the source pin.
    No source row is collapsed by accession or gene alias.
    """
    digest = hashlib.sha256(data).hexdigest()
    if digest != expected_sha256:
        raise ValueError("PXD036717 workbook SHA-256 mismatch")
    document = {
        "schemaVersion": 1, "accession": "PXD036717", "releaseEligible": False,
        "source": {"file": SOURCE_FILE, "sha256": digest, "bytes": len(data)},
        "method": {
            "quantity": "Reconstructed log2 APEX2/control LFQ enrichment",
            "inputs": "Deposited MaxQuant LFQ intensities, not raw mass spectra",
            "missing": "Null unless both intensities are finite and positive; no imputation",
            "replicates": "Source column indexing A/A, B/B, C/C; no paired-design inference",
            "mean": "Arithmetic mean of all three reconstructed log2 ratios; null otherwise",
            "comparisonAbsTolerance": ABS_TOLERANCE,
        },
        "sheets": [],
    }
    with ExitStack() as stack:
        original = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        stack.callback(original.close)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        stack.callback(cached.close)
        if original.sheetnames != [f"{bait}-APEX2" for bait in BAITS]:
            raise ValueError("PXD036717 workbook sheet names/order changed")
        for bait in BAITS:
            sheet = original[f"{bait}-APEX2"]
            value_sheet = cached[sheet.title]
            source_rows, cached_rows = iter(sheet), iter(value_sheet)
            header = next(source_rows)
            cached_header = next(cached_rows)
            expected = expected_headers(bait)
            if ([cell.value for cell in header[:32]] != expected
                    or [cell.value for cell in cached_header[:32]] != expected
                    or any(cell.value is not None for cell in header[32:])):
                raise ValueError(f"{sheet.title}: source headers changed")
            table = {
                "bait": bait, "sheet": sheet.title, "headers": expected,
                "dimensions": [sheet.max_row, sheet.max_column],
                "blankRows": [], "rows": [],
            }
            counts = Counter()
            ids = Counter()
            valid_by_replicate = Counter({rep: 0 for rep in REPLICATES})
            row_pairs = zip(source_rows, cached_rows, strict=True)
            for number, (source, values) in enumerate(row_pairs, 2):
                if any(cell.value is not None for cell in source[32:]):
                    raise ValueError(f"{sheet.title}:{number}: populated extra column")
                if all(cell.value is None for cell in source):
                    table["blankRows"].append(number)
                    continue
                identifier = source[0].value
                if (not isinstance(identifier, str) or not identifier.strip()
                        or source[0].data_type == "f"):
                    raise ValueError(f"{sheet.title}:A{number}: missing/nonliteral identifier")
                ids[identifier] += 1
                cells = [
                    _source_cell(cell, cache, f"{get_column_letter(i)}{number}")
                    for i, (cell, cache) in enumerate(zip(source[:32], values[:32], strict=True), 1)
                ]
                comparisons = []
                for i, rep in enumerate(REPLICATES):
                    indices = (2 + i, 5 + i, 8 + i, 11 + i, 14 + i)
                    comparison = reconstruct_comparison(*(values[j].value for j in indices))
                    comparison.update({
                        "replicate": rep,
                        "cells": dict(zip(
                            ("controlIntensity", "baitIntensity", "publishedControlLog2",
                             "publishedBaitLog2", "publishedRatio"),
                            (f"{get_column_letter(j + 1)}{number}" for j in indices), strict=True,
                        )),
                    })
                    counts.update(comparison["flags"])
                    valid_by_replicate[rep] += comparison["reconstructedLog2Ratio"] is not None
                    comparisons.append(comparison)
                ratios = [c["reconstructedLog2Ratio"] for c in comparisons]
                complete = all(value is not None for value in ratios)
                detection_counts = Counter({state: 0 for state in DETECTION_STATES})
                detection_counts.update(c["detectionState"] for c in comparisons)
                table["rows"].append({
                    "row": number, "accession": identifier, "cells": cells,
                    "comparisons": comparisons,
                    "validReplicates": sum(v is not None for v in ratios),
                    "detectionCounts": dict(detection_counts),
                    "meanLog2Ratio": math.fsum(ratios) / 3 if complete else None,
                })
            if not table["rows"]:
                raise ValueError(f"{sheet.title}: no data rows")
            table["audit"] = {
                "rowCount": len(table["rows"]), "validByReplicate": dict(valid_by_replicate),
                "completeMeans": sum(row["meanLog2Ratio"] is not None for row in table["rows"]),
                "flags": dict(sorted(counts.items())),
                "duplicateAccessions": {key: count for key, count in ids.items() if count > 1},
                "multiAccessionRows": [
                    row["row"] for row in table["rows"] if ";" in row["accession"]
                ],
                "formulaCells": sum(
                    cell["type"] == "f" for row in table["rows"] for cell in row["cells"]
                ),
                "detectionCounts": {
                    state: sum(row["detectionCounts"][state] for row in table["rows"])
                    for state in DETECTION_STATES
                },
            }
            document["sheets"].append(table)
    return document


def main(argv: list[str] | None = None) -> int:
    """Write non-release evidence JSON without changing source or browser files."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, required=True, help="Pinned mmc1.xlsx")
    parser.add_argument("--output", type=Path, required=True, help="Non-release evidence JSON")
    args = parser.parse_args(argv)
    if (args.input.resolve() == args.output.resolve()
            or (args.output.exists() and args.output.samefile(args.input))):
        parser.error("output must not overwrite the source workbook")
    document = read_workbook(args.input.read_bytes())
    text = json.dumps(document, indent=2, ensure_ascii=False, allow_nan=False) + "\n"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(text, encoding="utf-8")
    print(json.dumps({table["bait"]: table["audit"] for table in document["sheets"]}, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
