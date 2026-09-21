#!/usr/bin/env python3
"""Extract the gene-associated TSS rows from Tan et al. 2018 Table S1."""

from __future__ import annotations

import argparse
import csv
import hashlib
from pathlib import Path

from openpyxl import load_workbook


SOURCE_SHA256 = "098ecbd204cd1042a6edee1d2a500eaeca4efe034c503e2ade3db1c605e79b00"
SOURCE_COLUMNS = (
    "TSS_ID", "Chromosome", "Type", "Strand", "TSS_position", "Locus_tags",
    "Distances to start codon",
    "CT1_RawReads", "CT2_RawReads", "D1_RawReads", "D2_RawReads",
    "HL1_RawReads", "HL2_RawReads", "HT1_RawReads", "HT2_RawReads",
    "DK_vs_CT_log2FoldChange", "DK_vs_CT_padj",
    "HL_vs_CT_log2FoldChange", "HL_vs_CT_padj",
    "HT_vs_CT_log2FoldChange", "HT_vs_CT_padj",
)
OUTPUT_COLUMNS = (
    "tss_id", "replicon", "type", "strand", "position", "locus_tag",
    "source_start_distance_nt",
    "control_1", "control_2", "dark_1", "dark_2", "high_light_1",
    "high_light_2", "high_temperature_1", "high_temperature_2",
    "dark_log2fc", "dark_padj", "high_light_log2fc", "high_light_padj",
    "high_temperature_log2fc", "high_temperature_padj",
)


def source_hash(path: Path) -> str:
    """Hash a workbook without loading it all into memory."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def extract_rows(workbook: Path) -> list[tuple[object, ...]]:
    """Return all gTSS rows in a stable, narrow Table S1 column order."""
    observed_hash = source_hash(workbook)
    if observed_hash != SOURCE_SHA256:
        raise ValueError(f"Tan Table S1 workbook hash mismatch: {observed_hash}")
    sheet = load_workbook(workbook, read_only=True, data_only=True)["Table S1"]
    source = sheet.iter_rows(values_only=True)
    next(source)  # Descriptive title above the column names.
    header = next(source)
    indexes = {name: index for index, name in enumerate(header) if name is not None}
    missing = set(SOURCE_COLUMNS) - indexes.keys()
    if missing:
        raise ValueError(f"Tan Table S1 missing columns: {sorted(missing)}")
    rows = []
    seen = set()
    for raw in source:
        if not raw[indexes["TSS_ID"]]:
            continue
        if raw[indexes["Type"]] != "gTSS":
            continue
        tss_id = raw[indexes["TSS_ID"]]
        if tss_id in seen:
            raise ValueError(f"Duplicate Tan gTSS: {tss_id}")
        seen.add(tss_id)
        rows.append(tuple(raw[indexes[column]] for column in SOURCE_COLUMNS))
    if len(rows) != 2475:
        raise ValueError(f"Expected 2,475 gTSS rows, found {len(rows)}")
    return rows


def write_table(workbook: Path, output: Path) -> None:
    """Write the checked, attribution-preserving gene-linked subset as TSV."""
    rows = extract_rows(workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(OUTPUT_COLUMNS)
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("output", type=Path)
    options = parser.parse_args()
    write_table(options.workbook, options.output)


if __name__ == "__main__":
    main()
