#!/usr/bin/env python3
"""Extract non-gTSS regulatory evidence from Tan et al. 2018 Table S1."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path
from typing import Collection

from openpyxl import load_workbook


SOURCE_SHA256 = "098ecbd204cd1042a6edee1d2a500eaeca4efe034c503e2ade3db1c605e79b00"
DERIVED_SHA256 = "98a19729bf47940bf6832e3f08c3548fc4ddc7bb866176ab70e9c92c7c8d3dbe"
TYPE_COUNTS = {"aTSS": 1380, "iTSS": 724, "nTSS": 229}
CURRENT_LOCUS_PATTERN = re.compile(r"M744_RS\d{5}\Z")
SOURCE_COLUMNS = (
    "TSS_ID", "Chromosome", "Type", "Strand", "TSS_position", "Locus_tags",
    "Distances to start codon",
    "CT1_RawReads", "CT2_RawReads", "D1_RawReads", "D2_RawReads",
    "HL1_RawReads", "HL2_RawReads", "HT1_RawReads", "HT2_RawReads",
    "DK_vs_CT_log2FoldChange", "DK_vs_CT_padj",
    "HL_vs_CT_log2FoldChange", "HL_vs_CT_padj",
    "HT_vs_CT_log2FoldChange", "HT_vs_CT_padj",
)
OUTPUT_COLUMNS = (
    "tss_id", "replicon", "type", "strand", "position", "source_locus_tag",
    "source_start_distance_nt", "mapping_status", "mapping_reason",
    "mapped_locus_tag",
    "dark_log2fc", "dark_padj", "high_light_log2fc", "high_light_padj",
    "high_temperature_log2fc", "high_temperature_padj",
    "control_1", "control_2", "dark_1", "dark_2", "high_light_1",
    "high_light_2", "high_temperature_1", "high_temperature_2",
)


def source_hash(path: Path) -> str:
    """Hash a workbook without loading it all into memory."""
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1 << 20), b""):
            digest.update(block)
    return digest.hexdigest()


def current_loci_from_genes(path: Path) -> set[str]:
    """Read the exact current CDS identifiers from the plotted gene dataset."""
    genes = json.loads(path.read_text(encoding="utf-8"))
    loci = {gene["id"] for gene in genes}
    if len(loci) != len(genes) or not all(CURRENT_LOCUS_PATTERN.fullmatch(x) for x in loci):
        raise ValueError("Current genes must have unique M744_RS locus identifiers")
    return loci


def mapping_fields(source_locus: object, current_loci: Collection[str]) -> tuple[str, str, str]:
    """Return an explicit exact-ID mapping result without name inference."""
    if source_locus in (None, ""):
        return "unassociated", "source_locus_missing", ""
    locus = str(source_locus)
    if not CURRENT_LOCUS_PATTERN.fullmatch(locus):
        return "unmapped", "source_locus_not_current_id", ""
    if locus not in current_loci:
        return "unmapped", "source_locus_absent_from_current_cds", ""
    return "mapped", "exact_current_locus_id", locus


def extract_rows(
    workbook: Path,
    current_loci: Collection[str],
    *,
    expected_type_counts: dict[str, int] = TYPE_COUNTS,
) -> list[tuple[object, ...]]:
    """Return every non-gTSS row with an explicit current-CDS mapping result."""
    observed_hash = source_hash(workbook)
    if observed_hash != SOURCE_SHA256:
        raise ValueError(f"Tan Table S1 workbook hash mismatch: {observed_hash}")
    sheet = load_workbook(workbook, read_only=True, data_only=True)["Table S1"]
    source = sheet.iter_rows(values_only=True)
    next(source)  # Descriptive title above the column names.
    header = next(source)
    indexes = {name: index for index, name in enumerate(header) if name is not None}
    missing = set(SOURCE_COLUMNS) - indexes.keys()
    if missing:
        raise ValueError(f"Tan Table S1 missing columns: {sorted(missing)}")

    rows: list[tuple[object, ...]] = []
    seen: set[str] = set()
    type_counts: Counter[str] = Counter()
    for raw in source:
        tss_id = raw[indexes["TSS_ID"]]
        tss_type = raw[indexes["Type"]]
        if not tss_id or tss_type == "gTSS":
            continue
        if tss_type not in expected_type_counts:
            raise ValueError(f"Unexpected Tan TSS type: {tss_type}")
        if tss_id in seen:
            raise ValueError(f"Duplicate Tan non-gTSS: {tss_id}")
        seen.add(tss_id)
        strand = raw[indexes["Strand"]]
        position = raw[indexes["TSS_position"]]
        if tss_id != f"{tss_type}{strand}{position}":
            raise ValueError(f"Tan TSS identifier disagrees with row fields: {tss_id}")
        source_locus = raw[indexes["Locus_tags"]]
        mapping = mapping_fields(source_locus, current_loci)
        selected = tuple(raw[indexes[column]] for column in SOURCE_COLUMNS)
        rows.append(selected[:7] + mapping + selected[15:] + selected[7:15])
        type_counts[tss_type] += 1

    if type_counts != Counter(expected_type_counts):
        raise ValueError(
            f"Expected non-gTSS counts {expected_type_counts}, found {dict(type_counts)}"
        )
    return rows


def write_table(workbook: Path, genes: Path, output: Path) -> None:
    """Write the checked, attribution-preserving regulatory subset as TSV."""
    rows = extract_rows(workbook, current_loci_from_genes(genes))
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(OUTPUT_COLUMNS)
        writer.writerows(rows)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("genes", type=Path, help="Current plotted site/data/genes.json")
    parser.add_argument("output", type=Path)
    options = parser.parse_args()
    write_table(options.workbook, options.genes, options.output)


if __name__ == "__main__":
    main()
