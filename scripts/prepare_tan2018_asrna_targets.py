#!/usr/bin/env python3
"""Extract Tan 2018 Table S8 potential antisense targets as source claims."""

from __future__ import annotations

import argparse
import csv
import hashlib
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


WORKBOOK_SHA256 = "098ecbd204cd1042a6edee1d2a500eaeca4efe034c503e2ade3db1c605e79b00"
COUNTS = {"dark": 77, "high_light": 21, "high_temperature": 3}
MIN_ABS_LOG2FC = {"dark": 1.5, "high_light": 1.5, "high_temperature": 1.0}
COLUMNS = (
    "tss_id", "comparison", "source_selection_min_abs_log2fc", "strand",
    "atss_position", "atss_log2fc",
    "potential_target_locus", "potential_target_symbol", "gtss_position",
    "gtss_log2fc", "source_product",
)
HEADER = (
    "TSS_ID", "Strand", "aTSS_Position", "aTSS_log2FoldChange",
    "Potential_target_genes", "Potential_target_gene_symbols", "gTSS_Position",
    "gTSS_log2FoldChange", "Product",
)


def extract(workbook: Path, regulatory_tsv: Path) -> list[tuple[object, ...]]:
    """Keep every potential pair and its own comparison, without gene inference."""
    if hashlib.sha256(workbook.read_bytes()).hexdigest() != WORKBOOK_SHA256:
        raise ValueError("Tan 2018 workbook checksum mismatch")
    with regulatory_tsv.open(encoding="utf-8", newline="") as handle:
        sites = {row["tss_id"]: row for row in csv.DictReader(handle, delimiter="\t")}
    sheet = load_workbook(workbook, read_only=True, data_only=True)["Table S8"]
    source = sheet.iter_rows(values_only=True)
    next(source)
    if tuple(next(source)[:len(HEADER)]) != HEADER:
        raise ValueError("Tan Table S8 columns changed")
    rows = []
    seen = set()
    counts: Counter[str] = Counter()
    comparison = None
    for record in source:
        identifier = record[0]
        if not identifier:
            continue
        if isinstance(identifier, str) and "_vs_CT:" in identifier:
            comparison = {
                "DK": "dark", "HL": "high_light", "HT": "high_temperature",
            }.get(identifier[:2])
            if comparison is None:
                raise ValueError("Unexpected Tan Table S8 comparison")
            threshold = MIN_ABS_LOG2FC[comparison]
            terms = [
                (site, operator, float(value))
                for site, operator, value in re.findall(
                    r"([ag]TSS)_Log2FoldChange\s*(>=|<=)\s*(-?\d+(?:\.\d+)?)",
                    identifier,
                )
            ]
            expected = [
                ("aTSS", ">=", threshold), ("gTSS", "<=", -threshold),
                ("aTSS", "<=", -threshold), ("gTSS", ">=", threshold),
            ]
            if terms != expected:
                raise ValueError("Tan Table S8 selection rule changed")
            continue
        if comparison is None or identifier not in sites or not identifier.startswith("aTSS"):
            raise ValueError(f"Tan Table S8 has an unmatched antisense site: {identifier}")
        site = sites[identifier]
        if (str(record[4]) != site["source_locus_tag"]
                or str(record[2]) != site["position"]
                or str(record[1]) != site["strand"]):
            raise ValueError(f"Tan Table S8 target or coordinate changed: {identifier}")
        key = (identifier, comparison)
        if key in seen:
            raise ValueError(f"Duplicate Tan Table S8 pair: {key}")
        seen.add(key)
        rows.append((identifier, comparison, MIN_ABS_LOG2FC[comparison],
                     record[1], record[2], record[3],
                     record[4], record[5], record[6], record[7], record[8]))
        counts[comparison] += 1
    if counts != Counter(COUNTS):
        raise ValueError(f"Tan Table S8 comparison counts changed: {dict(counts)}")
    return rows


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("regulatory_tsv", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    rows = extract(args.workbook, args.regulatory_tsv)
    with args.output.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(COLUMNS)
        writer.writerows(rows)


if __name__ == "__main__":
    main()
